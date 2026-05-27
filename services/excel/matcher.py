import os
import pandas as pd
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.styles import Alignment
from core.excel_utils import ExcelFile, col_letter_to_index
from core.excel_utils import copy_merge_structure
from services.data.reader import DataReader

class ExcelMatcher:
    """Excel 数据匹配填充核心类（支持多列值填充）"""

    def __init__(self, target_path, source_path, config, sheet_indices=None, data_reader=None):
        self.target_path = target_path
        self.source_path = source_path
        self.config = config
        self.sheet_indices = sheet_indices
        self.errors = []
        self.target_wb = None
        self.target_ws = None
        self.source_df = None
        self.data_reader = data_reader or DataReader()
        self.allow_duplicate_keys = config.get('allow_duplicate_keys', False)

        # 解析配置
        self.target_key_cols = config['target_key_cols']          # 列字母列表，如 ['A', 'B']
        self.source_match_cols = config['source_match_cols']      # 列字母列表
        self.source_value_cols = config['source_value_cols']      # 列字母列表，如 ['B', 'C']
        self.result_cols = config['result_cols']                  # 结果列列表，可以是列字母或新列名

        # 验证长度一致性
        if len(self.source_value_cols) != len(self.result_cols):
            self.errors.append(f"源值列数量({len(self.source_value_cols)})与结果列数量({len(self.result_cols)})不一致")
        if len(self.target_key_cols) != len(self.source_match_cols):
            self.errors.append(f"目标主键列数量({len(self.target_key_cols)})与源匹配列数量({len(self.source_match_cols)})不一致")

    def _validate_column_letters(self):
        """验证列字母有效性，并转换索引"""
        try:
            # 转换目标主键列索引
            self.target_key_indices = [column_index_from_string(col) for col in self.target_key_cols]
            # 转换源匹配列索引
            self.source_match_indices = [col_letter_to_index(col) for col in self.source_match_cols]
            # 转换源值列索引
            self.source_value_indices = [col_letter_to_index(col) for col in self.source_value_cols]

            # 处理结果列：区分是列字母还是新列名
            self.result_is_new = []
            self.result_col_indices = []
            self.result_col_names = []
            for rc in self.result_cols:
                try:
                    idx = column_index_from_string(rc)
                    self.result_is_new.append(False)
                    self.result_col_indices.append(idx)
                    self.result_col_names.append(None)
                except:
                    self.result_is_new.append(True)
                    self.result_col_indices.append(None)
                    self.result_col_names.append(rc.strip())
        except Exception as e:
            self.errors.append(f"列字母解析错误: {str(e)}")
            return False
        return True

    def _load_source(self):
        """加载源文件（支持多sheet）"""
        try:
            if self.sheet_indices:
                all_sheets = pd.read_excel(self.source_path, sheet_name=None, dtype=str)
                sheet_names = list(all_sheets.keys())
                selected_sheets = []
                for idx in self.sheet_indices:
                    if idx < 1 or idx > len(sheet_names):
                        self.errors.append(f"Sheet索引 {idx} 超出范围（共 {len(sheet_names)} 个sheet）")
                        return False
                    selected_sheets.append(sheet_names[idx-1])
                df_list = []
                for sheet_name in selected_sheets:
                    df = all_sheets[sheet_name]
                    df_list.append(df)
                self.source_df = pd.concat(df_list, ignore_index=True, sort=False)
            else:
                self.source_df = self.data_reader.read_excel(self.source_path, dtype=str)

            if self.source_df.empty:
                self.errors.append("源文件无数据")
                return False

            self.source_col_names = list(self.source_df.columns)
            max_cols = len(self.source_col_names)

            # 验证列范围
            for col_letter, idx in zip(self.source_match_cols, self.source_match_indices):
                if idx > max_cols:
                    self.errors.append(f"源文件匹配列 {col_letter} 超出列范围（共 {max_cols} 列）")
                    return False
            for col_letter, idx in zip(self.source_value_cols, self.source_value_indices):
                if idx > max_cols:
                    self.errors.append(f"源文件值列 {col_letter} 超出列范围（共 {max_cols} 列）")
                    return False
            return True
        except Exception as e:
            self.errors.append(f"读取源文件失败: {str(e)}")
            return False

    def _load_target(self):
        """加载目标文件"""
        try:
            self.target_wb = ExcelFile(self.target_path)
            self.target_wb.load()
            self.target_ws = self.target_wb.active_sheet
            max_col = self.target_ws.max_column
            for col_letter, idx in zip(self.target_key_cols, self.target_key_indices):
                if idx > max_col:
                    self.errors.append(f"目标文件键列 {col_letter} 超出最大列数 {max_col}")
                    return False
            return True
        except Exception as e:
            msg = str(e)
            if 'does not support' in msg or 'file format' in msg:
                self.errors.append("目标文件格式不支持，请使用 .xlsx 格式文件。如为 .xls 文件，请先用 Excel 另存为 .xlsx。")
            else:
                self.errors.append(f"读取目标文件失败: {msg}")
            return False

    def _extract_target_keys(self):
        """提取目标文件中所有行的组合主键"""
        start_row = 2
        end_row = self.target_ws.max_row
        keys = []
        for row in range(start_row, end_row + 1):
            key = []
            for col in self.target_key_indices:
                val = self.target_wb.get_cell_value(row, col, sheet=self.target_ws)
                key.append(str(val).strip() if val else '')
            keys.append(tuple(key))
        return keys, start_row, end_row

    def _build_lookup_dict(self):
        """构建源文件的查找字典：key -> 元组(值列1, 值列2, ...)"""
        # 获取列名
        match_col_names = [self.source_col_names[idx - 1] for idx in self.source_match_indices]
        value_col_names = [self.source_col_names[idx - 1] for idx in self.source_value_indices]

        lookup = {}
        for _, row in self.source_df.iterrows():
            key = tuple(str(row[col]).strip() if pd.notna(row[col]) else '' for col in match_col_names)
            values = tuple(
                row[col] if pd.notna(row[col]) else 0
                for col in value_col_names
            )
            # 如果允许重复键，保留第一次出现（与原行为一致，不覆盖）
            if key not in lookup:
                lookup[key] = values
        return lookup

    def _check_duplicate_target_keys(self, target_keys):
        """检查目标键是否重复，返回重复的键列表"""
        seen = set()
        duplicates = set()
        for key in target_keys:
            if key in seen:
                duplicates.add(key)
            else:
                seen.add(key)
        return duplicates

    def _write_result_columns(self, keys, lookup, start_row):
        """写入多个结果列（使用 self.result_col_indices 并原地更新新建列的索引）"""
        # 先处理新建列（插入列），更新 self.result_col_indices
        for i, (is_new, idx, col_name) in enumerate(zip(self.result_is_new, self.result_col_indices, self.result_col_names)):
            if is_new:
                max_col = self.target_ws.max_column
                new_idx = max_col + 1
                self.result_col_indices[i] = new_idx
                self.target_ws.insert_cols(new_idx)
                self.target_ws.cell(1, new_idx).value = col_name

        # 清除每个结果列的合并单元格
        for idx in self.result_col_indices:
            to_remove = []
            for merged in self.target_ws.merged_cells.ranges:
                if merged.min_col <= idx <= merged.max_col:
                    to_remove.append(merged)
            for merged in to_remove:
                self.target_ws.unmerge_cells(str(merged))

        # 写入数据
        seen_keys = set()
        for row_offset, key in enumerate(keys):
            row_num = start_row + row_offset
            if self.allow_duplicate_keys:
                if key in seen_keys:
                    continue
                seen_keys.add(key)

            values_tuple = lookup.get(key, tuple(0 for _ in self.result_col_indices))
            for i, idx in enumerate(self.result_col_indices):
                val = values_tuple[i] if i < len(values_tuple) else 0
                # 尝试转数值
                try:
                    val = float(val)
                except (TypeError, ValueError):
                    pass
                self.target_ws.cell(row_num, idx).value = val

        # 复制合并单元格结构（基于第一个主键列）
        if self.target_key_indices:
            ref_col = self.target_key_indices[0]
            for idx in self.result_col_indices:
                copy_merge_structure(self.target_ws, ref_col, self.target_ws, idx, start_row, self.target_ws.max_row)

        # 设置右对齐
        for idx in self.result_col_indices:
            for row in range(start_row, self.target_ws.max_row + 1):
                cell = self.target_ws.cell(row, idx)
                cell.alignment = Alignment(horizontal='right')

    def _generate_report(self, used_target_keys, lookup, start_row):
        """生成校验报告 sheet，展示所有源值列的值"""
        target_keys_full, _, _ = self._extract_target_keys()
        target_keys_used = used_target_keys
        source_keys_set = set(lookup.keys())

        # 准备列标题：行号、目标键、源匹配键、每个源值列的值、匹配状态、备注
        header_row = 1
        target_col_names = []
        for col_1based in self.target_key_indices:
            col_name = self.target_ws.cell(header_row, col_1based).value
            if not col_name:
                col_name = f"列{get_column_letter(col_1based)}"
            target_col_names.append(col_name)
        target_key_label = "|".join(target_col_names)
        target_file_name = os.path.basename(self.target_path)

        # 源值列列名（使用列字母作为标识）
        source_value_labels = [f"源文件值({col})" for col in self.source_value_cols]

        report_sheet = self.target_wb._workbook.create_sheet("校验报告")
        headers = ["行号", f"目标文件键列值\n({target_file_name}【{target_key_label}】)", "源文件匹配键"] + source_value_labels + ["匹配状态", "备注"]
        for col, header in enumerate(headers, 1):
            report_sheet.cell(1, col).value = header

        row_idx = 2
        # 遍历目标文件每一行（包括未匹配的）
        for i, (full_key, used_key) in enumerate(zip(target_keys_full, target_keys_used)):
            row_num = start_row + i
            if used_key in source_keys_set:
                val_tuple = lookup.get(used_key, tuple())
                # 匹配成功
                row_data = [row_num, "|".join(str(v) for v in full_key), "|".join(str(v) for v in used_key)]
                # 添加每个值列的值
                for col_idx in range(len(self.source_value_cols)):
                    val = val_tuple[col_idx] if col_idx < len(val_tuple) else ''
                    row_data.append(val)
                row_data.append(1)
                row_data.append('')
                for col_idx, val in enumerate(row_data, 1):
                    report_sheet.cell(row_idx, col_idx).value = val
                row_idx += 1
            else:
                # 仅目标存在，未匹配
                row_data = [row_num, "|".join(str(v) for v in full_key), "|".join(str(v) for v in used_key)]
                # 值列留空
                for _ in self.source_value_cols:
                    row_data.append('')
                row_data.append(0)
                row_data.append('')
                for col_idx, val in enumerate(row_data, 1):
                    report_sheet.cell(row_idx, col_idx).value = val
                row_idx += 1

        # 仅源文件存在的记录
        for key, val_tuple in lookup.items():
            if key not in set(target_keys_used):
                row_data = ['-', '', "|".join(str(v) for v in key)]
                for col_idx in range(len(self.source_value_cols)):
                    val = val_tuple[col_idx] if col_idx < len(val_tuple) else ''
                    row_data.append(val)
                row_data.append(0)
                row_data.append('仅源文件')
                for col_idx, val in enumerate(row_data, 1):
                    report_sheet.cell(row_idx, col_idx).value = val
                row_idx += 1

        # 设置对齐
        for col in range(1, len(headers) + 1):
            for row in range(2, row_idx):
                cell = report_sheet.cell(row, col)
                cell.alignment = Alignment(horizontal='right' if col > 1 else 'left')

    def match(self):
        if not self._validate_column_letters():
            return False, self.errors
        if not self._load_source():
            return False, self.errors
        if not self._load_target():
            return False, self.errors

        target_keys, start_row, end_row = self._extract_target_keys()

        if not self.allow_duplicate_keys:
            duplicates = self._check_duplicate_target_keys(target_keys)
            if duplicates:
                dup_examples = [f"'{k}'" for k in list(duplicates)[:5]]
                self.errors.append(f"目标文件键列存在重复值: {', '.join(dup_examples)}。如需继续填充，请勾选“允许目标键重复”。")
                return False, self.errors

        lookup = self._build_lookup_dict()
        match_len = len(self.target_key_cols)
        used_target_keys = [key[:match_len] for key in target_keys]

        # 写入结果列
        self._write_result_columns(used_target_keys, lookup, start_row)

        # 生成报告（多列版本）
        self._generate_report(used_target_keys, lookup, start_row)

        out_path = os.path.join(os.path.dirname(self.target_path), f"matched_{os.path.basename(self.target_path)}")
        self.target_wb.save(out_path)
        return True, out_path

    def debug_info(self):
        if not self._validate_column_letters():
            return {'error': self.errors}
        if not self._load_source():
            return {'error': self.errors}
        if not self._load_target():
            return {'error': self.errors}

        target_keys, start_row, end_row = self._extract_target_keys()
        lookup = self._build_lookup_dict()
        match_len = len(self.target_key_cols)
        used_target_keys = [key[:match_len] for key in target_keys]

        target_set = set(used_target_keys)
        source_set = set(lookup.keys())
        common = target_set & source_set
        only_target = target_set - source_set
        only_source = source_set - target_set

        matched_pairs = []
        unmatched_pairs = []
        for i, (full_key, used_key) in enumerate(zip(target_keys, used_target_keys)):
            row_num = start_row + i
            matched_val = lookup.get(used_key)
            if matched_val is not None:
                matched_pairs.append((row_num, full_key, used_key, matched_val))
            else:
                unmatched_pairs.append((row_num, full_key, used_key))
            if len(matched_pairs) >= 20 and len(unmatched_pairs) >= 20:
                break

        return {
            'config': self.config,
            'target_rows': len(target_keys),
            'target_keys_sample': [(start_row + i, key) for i, key in enumerate(target_keys[:20])],
            'source_keys_sample': list(source_set)[:20],
            'source_values_sample': [(k, lookup[k]) for k in list(source_set)[:20]],
            'common_count': len(common),
            'only_target_count': len(only_target),
            'only_source_count': len(only_source),
            'matched_count': len(common),
            'unmatched_keys': list(only_target)[:20],
            'start_row': start_row,
            'end_row': end_row,
            'matched_pairs': matched_pairs,
            'unmatched_pairs': unmatched_pairs
        }
