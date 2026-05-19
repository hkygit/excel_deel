# services/excel/calculator.py
import numpy as np
from typing import List, Tuple
from openpyxl.utils import column_index_from_string
from openpyxl.styles import Alignment
from core.excel_utils import ExcelFile
from core.expression_engine import get_expression_engine, extract_column_letters

class ColumnCalculator:
    def __init__(self, excel_file: ExcelFile):
        self.excel = excel_file
        self.ws = excel_file.active_sheet

    def get_data_range(self, header_row: int = 1) -> Tuple[int, int]:
        return header_row + 1, self.ws.max_row

    def calculate_column(self, expression: str, result_col: str,
                         merge_by_cols: List[str] = None) -> None:
        engine = get_expression_engine()
        col_letters = extract_column_letters(expression)
        if not col_letters:
            raise ValueError("表达式中未包含列字母")

        try:
            result_col_idx = column_index_from_string(result_col)
            is_new_col = False
        except:
            is_new_col = True
            result_col_name = result_col.strip()

        if not is_new_col:
            result_col_idx = column_index_from_string(result_col)
        else:
            max_col = self.ws.max_column
            result_col_idx = max_col + 1
            self.ws.insert_cols(result_col_idx)
            self.ws.cell(1, result_col_idx).value = result_col_name

        start_row, end_row = self.get_data_range()

        col_data = {}
        for col_letter in col_letters:
            col_idx = column_index_from_string(col_letter)
            vals = []
            for row in range(start_row, end_row + 1):
                val = self.excel.get_cell_value(row, col_idx, sheet=self.ws)
                try:
                    val = float(val) if val is not None else 0
                except:
                    val = 0
                vals.append(val)
            col_data[col_letter] = vals

        calc_vals = []
        first_letter = list(col_data.keys())[0]
        for i in range(len(col_data[first_letter])):
            row_vals = {letter: col_data[letter][i] for letter in col_letters}
            result = engine.evaluate(expression, row_vals)
            calc_vals.append(result)

        self._clear_merges_in_column(result_col_idx)
        for i, val in enumerate(calc_vals):
            row = start_row + i
            self.ws.cell(row, result_col_idx).value = None if np.isnan(val) else val

        ref_col = sorted(col_letters)[0]
        ref_idx = column_index_from_string(ref_col)
        self.excel.copy_merge_structure(ref_idx, result_col_idx, start_row, end_row, sheet=self.ws)

        if merge_by_cols:
            self._merge_by_keys(merge_by_cols, result_col_idx, start_row, end_row)

        for row in range(start_row, end_row + 1):
            self.ws.cell(row, result_col_idx).alignment = Alignment(horizontal='right')

    def _clear_merges_in_column(self, col_idx: int):
        to_remove = []
        for merged in self.ws.merged_cells.ranges:
            if merged.min_col <= col_idx <= merged.max_col:
                to_remove.append(merged)
        for merged in to_remove:
            self.ws.unmerge_cells(str(merged))

    def _merge_by_keys(self, merge_cols, result_col_idx, start_row, end_row):
        merge_indices = [column_index_from_string(c.strip()) for c in merge_cols]
        group_first_row = {}
        for row in range(start_row, end_row + 1):
            key = tuple(self.excel.get_cell_value(row, col, sheet=self.ws) for col in merge_indices)
            if key not in group_first_row:
                group_first_row[key] = row

        for row in range(start_row, end_row + 1):
            key = tuple(self.excel.get_cell_value(row, col, sheet=self.ws) for col in merge_indices)
            first_row = group_first_row.get(key)
            if row != first_row:
                self.ws.cell(row, result_col_idx).value = None

        for key, first_row in group_first_row.items():
            last_row = first_row
            for row in range(first_row, end_row + 1):
                if tuple(self.excel.get_cell_value(row, col, sheet=self.ws) for col in merge_indices) == key:
                    last_row = row
                else:
                    break
            if first_row < last_row:
                self.ws.merge_cells(start_row=first_row, start_column=result_col_idx,
                                    end_row=last_row, end_column=result_col_idx)
