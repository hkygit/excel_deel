# services/convert/service.py
import os, tempfile, zipfile, csv, logging
import pandas as pd
from typing import Optional, Callable
from core.exceptions import ValidationError
from core.task.base import Task

logger = logging.getLogger(__name__)

class ConvertService(Task):
    def __init__(self, file_path: str, target_format: str = 'csv',
                 encoding: str = 'utf-8-sig', separator: str = ',',
                 progress_callback: Optional[Callable] = None):
        self.file_path = file_path
        self.target_format = target_format
        self.encoding = encoding
        self.separator = separator
        self.progress_callback = progress_callback

    def precheck(self) -> dict:
        ext = os.path.splitext(self.file_path)[1].lower()
        if ext in ('.xlsx', '.xls', '.xlsm'):
            from openpyxl import load_workbook
            wb = load_workbook(self.file_path, read_only=True)
            rows = sum(ws.max_row - 1 for ws in wb.worksheets)
            wb.close()
            return {"type": "excel", "sheets": len(wb.sheetnames), "total_rows": rows}
        else:
            size = os.path.getsize(self.file_path)
            return {"type": "csv", "total_rows": max(size // 200, 1)}

    def execute(self) -> str:
        ext = os.path.splitext(self.file_path)[1].lower()
        if self.target_format == 'csv' and ext in ('.xlsx', '.xls', '.xlsm'):
            return self._excel_to_csv()
        elif self.target_format == 'xlsx' and ext in ('.csv',):
            return self._csv_to_excel()
        else:
            raise ValidationError(f"不支持的转换: {ext} -> {self.target_format}")

    def _excel_to_csv(self) -> str:
        from openpyxl import load_workbook
        wb = load_workbook(self.file_path, read_only=True)
        sheet_names = wb.sheetnames
        if len(sheet_names) == 1:
            ws = wb.active
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            total_rows = ws.max_row - 1
            with tempfile.NamedTemporaryFile(suffix='.csv', delete=False, mode='w', encoding=self.encoding, newline='') as f:
                writer = csv.writer(f, delimiter=self.separator)
                writer.writerow(headers)
                buffer = []
                buffer_size = 200000   # 20 万行内存缓冲
                processed = 0
                self._update_progress(0, f"开始转换 (共约 {total_rows:,} 行)...")
                for row in ws.iter_rows(min_row=2, values_only=True):
                    buffer.append(row)
                    processed += 1
                    if len(buffer) >= buffer_size:
                        writer.writerows(buffer)
                        buffer.clear()
                        # 只在大块完成时更新进度
                        pct = int(processed / total_rows * 100) if total_rows else 50
                        self._update_progress(pct, f"已转换 {processed:,} / {total_rows:,} 行")
                # 写入剩余行
                if buffer:
                    writer.writerows(buffer)
                out_path = f.name
            wb.close()
            self._update_progress(100, "转换完成")
            return out_path
        else:
            # 多 sheet 打包不变
            zip_path = tempfile.NamedTemporaryFile(suffix='.zip', delete=False).name
            tmp_files = []
            total = len(sheet_names)
            with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
                for i, name in enumerate(sheet_names):
                    self._update_progress(i/total*100, f"转换工作表 {name} ({i+1}/{total})")
                    df = pd.read_excel(self.file_path, sheet_name=name, engine='openpyxl')
                    csv_name = f"{os.path.splitext(os.path.basename(self.file_path))[0]}_{name}.csv"
                    tmp_csv = tempfile.NamedTemporaryFile(suffix='.csv', delete=False, mode='w', encoding=self.encoding)
                    df.to_csv(tmp_csv.name, index=False, encoding=self.encoding, sep=self.separator)
                    tmp_files.append(tmp_csv.name)
                    zf.write(tmp_csv.name, arcname=csv_name)
            for f in tmp_files:
                try: os.unlink(f)
                except: pass
            wb.close()
            self._update_progress(100, "转换完成")
            return zip_path

    def _csv_to_excel(self) -> str:
        total_size = os.path.getsize(self.file_path)
        estimated_rows = max(total_size // 100, 1)
        self._update_progress(0, "开始读取CSV...")

        chunks = []
        total_rows = 0
        for chunk in pd.read_csv(self.file_path, encoding=self.encoding, sep=self.separator, chunksize=100000):
            chunks.append(chunk)
            total_rows += len(chunk)
            pct = min(int(total_rows / estimated_rows * 100), 99)
            self._update_progress(pct, f"已读取 {total_rows:,} 行")

        df = pd.concat(chunks)
        self._update_progress(95, "正在写入Excel...")
        out_path = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False).name
        df.to_excel(out_path, index=False, engine='openpyxl')
        self._update_progress(100, "转换完成")
        return out_path

    def _update_progress(self, percent, message):
        if self.progress_callback:
            self.progress_callback(percent=percent, message=message)
