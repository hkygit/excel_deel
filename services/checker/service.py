import os
import time
import logging
from typing import List, Dict, Optional, Callable, Set
import pandas as pd
from openpyxl import Workbook

from core.progress_task import ProgressTask
from core.exceptions import ValidationError
from services.data.stream_reader import StreamReader
from services.data.row_counter import RowCounter

logger = logging.getLogger(__name__)

BATCH_SIZE = 10000

class CheckerService(ProgressTask):
    """文件校验服务（优化版：内存优先，分批流式对比）"""

    def __init__(self, files_a=None, files_b=None,
                 key_map_a=None, key_map_b=None,
                 compare_map=None, filters_a=None,
                 duplicate_action=None, config=None,
                 progress_callback=None,
                 precheck_data=None):
        super().__init__(progress_callback)
        self.files_a = files_a or []
        self.files_b = files_b or []
        self.key_map_a = key_map_a or []
        self.key_map_b = key_map_b or []
        self.compare_map = compare_map or []
        self.filters_a = filters_a or []
        self.config = config
        self.duplicate_action = duplicate_action or (config.duplicate_action if config else "error")
        self.precheck_data = precheck_data

    def precheck(self) -> dict:
        ta = sum(RowCounter.count(f) for f in self.files_a)
        tb = sum(RowCounter.count(f) for f in self.files_b)
        total = ta + tb
        total_size = sum(os.path.getsize(f) for f in self.files_a + self.files_b) / (1024 * 1024)
        return {
            "total_files_a": len(self.files_a),
            "total_files_b": len(self.files_b),
            "total_rows_a": ta, "total_rows_b": tb,
            "total_rows": total,
            "estimated_memory_mb": round(total * 200 / (1024*1024), 2),
            "estimated_time_seconds": round(total / 20000, 2),
            "total_size_mb": round(total_size, 2),
            "advice": "数据量大" if total > 500000 else "安全",
            "can_proceed": True
        }

    @staticmethod
    def _make_key(row, key_cols):
        """统一构造主键（字符串元组），避免 numpy 类型问题"""
        return tuple(str(row[k]) for k in key_cols)

    def execute(self) -> str:
        start_time = time.time()
        if self.precheck_data:
            total_rows_a = self.precheck_data.get('total_rows_a', 0)
            total_rows_b = self.precheck_data.get('total_rows_b', 0)
        else:
            total_rows_a = sum(RowCounter.count(f) for f in self.files_a)
            total_rows_b = sum(RowCounter.count(f) for f in self.files_b)
        self._total_rows = total_rows_a + total_rows_b

        use_sqlite_fallback = False
        max_memory_rows = self.config.max_memory_rows if self.config else 200000
        if total_rows_a > max_memory_rows:
            use_sqlite_fallback = True
            logger.info(f"标准群行数 {total_rows_a} 超过阈值 {max_memory_rows}，启用 SQLite 兜底模式")

        if use_sqlite_fallback:
            return self._execute_sqlite_fallback(start_time)
        else:
            return self._execute_memory_stream(start_time, total_rows_a, total_rows_b)

    # ------------------------------------------------------------------
    # 内存 + 流式对比
    # ------------------------------------------------------------------
    def _execute_memory_stream(self, start_time, total_rows_a, total_rows_b):
        self.set_progress(0, "读取标准群...")

        needed_a = set(self.key_map_a)
        for m in self.compare_map:
            needed_a.add(m["standard"])
        for f in self.filters_a:
            if "field" in f:
                needed_a.add(f["field"])
        needed_a.discard("")

        std_rows = []
        processed_a = 0
        for file in self.files_a:
            for row in StreamReader.read_excel(file, list(needed_a), self.filters_a):
                std_rows.append(row)
                processed_a += 1
                if processed_a % 5000 == 0:
                    pct = min(int(processed_a / total_rows_a * 100), 30) if total_rows_a else 0
                    self.set_progress(pct, f"读取标准群 {processed_a}/{total_rows_a}")
        std_df = pd.DataFrame(std_rows)
        if std_df.empty:
            raise ValidationError("标准群无有效数据")

        if self.duplicate_action == "error":
            if std_df.duplicated(subset=self.key_map_a).any():
                raise ValidationError("标准群主键重复")
        else:
            std_df = std_df.drop_duplicates(subset=self.key_map_a, keep='first')

        std_df = std_df.set_index(self.key_map_a)

        cmp_cols_a = [m["standard"] for m in self.compare_map]
        cmp_cols_b = [m["compare"] for m in self.compare_map]

        # 流式写入 Excel
        wb = Workbook(write_only=True)
        ws_diff = wb.create_sheet("差异数据")
        ws_std_only = wb.create_sheet("标准群独有")
        ws_cmp_only = wb.create_sheet("比较群独有")

        diff_header = (list(self.key_map_a) +
                       [f"{col}_标准" for col in cmp_cols_a] +
                       [f"{col}_比较" for col in cmp_cols_b])
        ws_diff.append(diff_header)
        ws_std_only.append(list(self.key_map_a) + cmp_cols_a)
        ws_cmp_only.append(list(self.key_map_b) + cmp_cols_b)

        matched_std_keys: Set[tuple] = set()

        self.set_progress(30, "开始流式对比...")

        needed_b = set(self.key_map_b)
        for m in self.compare_map:
            needed_b.add(m["compare"])
        needed_b.discard("")

        batch_buffer = []
        processed_b = 0
        cmp_only_rows = []

        for file in self.files_b:
            for row in StreamReader.read_excel(file, list(needed_b)):
                batch_buffer.append(row)
                processed_b += 1
                if len(batch_buffer) >= BATCH_SIZE:
                    self._process_batch(batch_buffer, std_df, cmp_cols_a, cmp_cols_b,
                                        ws_diff, matched_std_keys, cmp_only_rows)
                    batch_buffer = []
                    pct = 30 + int((processed_b / total_rows_b) * 65) if total_rows_b else 50
                    self.set_progress(pct, f"对比中 {processed_b}/{total_rows_b}")

        if batch_buffer:
            self._process_batch(batch_buffer, std_df, cmp_cols_a, cmp_cols_b,
                                ws_diff, matched_std_keys, cmp_only_rows)

        for row in cmp_only_rows:
            ws_cmp_only.append(list(row))

        # 标准群独有（统一用字符串主键）
        self.set_progress(96, "计算标准群独有...")
        all_std_keys = set(self._make_key(row, self.key_map_a) for row in std_df.reset_index().to_dict('records'))
        only_std_keys = all_std_keys - matched_std_keys
        if only_std_keys:
            std_only_df = std_df.reset_index()
            std_only_df['_key'] = std_only_df.apply(lambda r: self._make_key(r, self.key_map_a), axis=1)
            for _, row in std_only_df[std_only_df['_key'].isin(only_std_keys)].iterrows():
                values = [row[k] for k in self.key_map_a] + [row[col] for col in cmp_cols_a]
                ws_std_only.append(values)

        import tempfile
        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        tmp.close()
        out_path = tmp.name
        wb.save(out_path)
        wb.close()

        self.set_progress(100, "完成")
        return out_path

    def _process_batch(self, batch_rows, std_df, cmp_cols_a, cmp_cols_b,
                       ws_diff, matched_std_keys, cmp_only_buf):
        cmp_df = pd.DataFrame(batch_rows)
        cmp_df['_key'] = cmp_df.apply(lambda r: self._make_key(r, self.key_map_b), axis=1)

        merged = pd.merge(
            std_df.reset_index(), cmp_df,
            left_on=self.key_map_a, right_on=self.key_map_b,
            how='inner', suffixes=('_标准', '_比较')
        )

        for _, match_row in merged.iterrows():
            matched_std_keys.add(self._make_key(match_row, self.key_map_a))

        if merged.empty:
            for _, row in cmp_df.iterrows():
                cmp_only_buf.append(row.drop(labels=['_key']).to_dict())
            return

        diff_mask = pd.Series(False, index=merged.index)
        for ca, cb in zip(cmp_cols_a, cmp_cols_b):
            diff_mask |= (merged[f'{ca}_标准'] != merged[f'{cb}_比较'])
        diff_rows = merged[diff_mask]

        if not diff_rows.empty:
            for _, row in diff_rows.iterrows():
                values = [row[k] for k in self.key_map_a]
                for ca in cmp_cols_a:
                    values.append(row[f'{ca}_标准'])
                for cb in cmp_cols_b:
                    values.append(row[f'{cb}_比较'])
                ws_diff.append(values)

        merged_keys = set(self._make_key(row, self.key_map_b) for _, row in merged.iterrows())
        for _, row in cmp_df.iterrows():
            if row['_key'] not in merged_keys:
                cmp_only_buf.append(row.drop(labels=['_key']).to_dict())

    # ------------------------------------------------------------------
    # SQLite 兜底模式（略作优化，与原版一致）
    # ------------------------------------------------------------------
    def _execute_sqlite_fallback(self, start_time: float) -> str:
        from core.sqlite_buffer import SqliteBuffer
        from services.data.report_writer import ReportWriter

        self.set_progress(0, "正在准备（SQLite模式）...")

        needed_a = set(self.key_map_a)
        for m in self.compare_map:
            needed_a.add(m["standard"])
        for f in self.filters_a:
            if "field" in f:
                needed_a.add(f["field"])
        needed_a.discard("")
        needed_b = set(self.key_map_b)
        for m in self.compare_map:
            needed_b.add(m["compare"])

        db = SqliteBuffer()
        try:
            db.create_table("std", list(needed_a), self.key_map_a)
            db.create_table("cmp", list(needed_b), self.key_map_b)

            self.set_progress(0, "导入标准群...")
            self._import_files(db, "std", self.files_a, list(needed_a), self.filters_a)
            self.set_progress(50, "导入比较群...")
            self._import_files(db, "cmp", self.files_b, list(needed_b), None)

            if self.duplicate_action == "error":
                for tbl, keys, label in [("std", self.key_map_a, "标准"), ("cmp", self.key_map_b, "比较")]:
                    sql = f"SELECT {','.join(keys)} FROM {tbl} GROUP BY {','.join(keys)} HAVING COUNT(*)>1"
                    if db.execute(sql).fetchone():
                        raise ValidationError(f"{label}群主键重复")
            else:
                db.remove_duplicates("std", self.key_map_a)
                db.remove_duplicates("cmp", self.key_map_b)

            cmp_a = [m["standard"] for m in self.compare_map]
            cmp_b = [m["compare"] for m in self.compare_map]
            join_on = " AND ".join([f"std.{ka}=cmp.{kb}" for ka,kb in zip(self.key_map_a, self.key_map_b)])
            eq_parts = [f'TRIM(LOWER(std.{ca})) = TRIM(LOWER(cmp.{cb}))' for ca,cb in zip(cmp_a,cmp_b)]
            diff_sql = f"SELECT std.*,cmp.* FROM std JOIN cmp ON {join_on} WHERE NOT ({' AND '.join(eq_parts)})"
            only_a_sql = f"SELECT std.* FROM std LEFT JOIN cmp ON {join_on} WHERE cmp.{self.key_map_b[0]} IS NULL"
            only_b_sql = f"SELECT cmp.* FROM cmp LEFT JOIN std ON {join_on} WHERE std.{self.key_map_a[0]} IS NULL"

            self.set_progress(80, "执行对比查询...")
            diff_df = db.fetch_df(diff_sql)
            only_a_df = db.fetch_df(only_a_sql)
            only_b_df = db.fetch_df(only_b_sql)

            self.set_progress(90, "写入结果...")
            sheets = {}
            if not diff_df.empty:
                sheets["差异数据"] = diff_df
            if not only_a_df.empty:
                sheets["标准群独有"] = only_a_df
            if not only_b_df.empty:
                sheets["比较群独有"] = only_b_df

            writer = ReportWriter(
                timeout_seconds=self.config.sheet2_timeout_seconds if self.config else None,
                progress_callback=lambda **kw: self.set_progress(**kw)
            )
            result_path = writer.write(sheets, start_time)
            self.set_progress(100, "完成")
            return result_path
        finally:
            db.close()

    def _import_files(self, db, table, files, columns, filters=None):
        batch = []
        batch_size = 5000
        total = self._total_rows if self._total_rows else 1
        for file in files:
            for row in StreamReader.read_excel(file, columns, filters):
                batch.append([row[c] for c in columns])
                self._processed_rows += 1
                if len(batch) >= batch_size:
                    db.insert_batch(table, columns, batch)
                    batch.clear()
            if batch:
                db.insert_batch(table, columns, batch)
                batch.clear()