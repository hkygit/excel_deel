# services/subtotal/engine.py
import os
import openpyxl
from openpyxl.styles import PatternFill, Font, Border, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from core.csv_utils import split_fields
from core.header_detection import SimpleExcelHeaderDetector
from core.file_naming import generate_output_filename

SUBTOTAL_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
SUBTOTAL_FONT = Font(bold=True)
THIN_BORDER = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))

def _build_header_map(ws):
    """根据工作表第一行构建列名 -> 列字母的映射"""
    header_map = {}
    for cell in ws[1]:
        if cell.value is not None:
            col_letter = get_column_letter(cell.column)
            header_map[str(cell.value).strip()] = col_letter
    return header_map

def resolve_single_col(token, header_map):
    """解析单个列标识（字母或列名），返回1-based列索引"""
    try:
        return column_index_from_string(token)
    except (ValueError, KeyError):
        if token in header_map:
            return column_index_from_string(header_map[token])
        raise ValueError(f"无效的列标识: {token}")

def parse_cols_with_header(raw, header_map):
    """解析列字符串（支持范围、列名），返回列索引列表"""
    if not raw:
        return []
    items = split_fields(raw)
    result = []
    for item in items:
        if '-' in item:
            start_str, end_str = item.split('-', 1)
            start_idx = resolve_single_col(start_str.strip(), header_map)
            end_idx = resolve_single_col(end_str.strip(), header_map)
            for i in range(start_idx, end_idx + 1):
                result.append(i)
        else:
            result.append(resolve_single_col(item.strip(), header_map))
    return sorted(list(set(result)))

class SubtotalEngine:
    def __init__(self, filepath, key_cols_str, calc_cols_str, agg_funcs,
                 dimension_cols=None, sheets='1', original_filename='data.xlsx'):
        self.filepath = filepath
        self.original_filename = original_filename
        self.key_cols_raw = key_cols_str
        self.calc_cols_raw = calc_cols_str
        self.dimension_cols_raw = dimension_cols or ''
        self.agg_funcs = agg_funcs
        self.sheets = sheets
        self.detector = SimpleExcelHeaderDetector()
        self.key_cols = []
        self.calc_cols = []
        self.dimension_cols = []

    def _get_agg_func(self, col_letter):
        return self.agg_funcs.get(col_letter, self.agg_funcs.get('default', 'sum'))

    def process(self):
        wb = openpyxl.load_workbook(self.filepath)
        first_ws = wb.active
        header_map = _build_header_map(first_ws)

        self.key_cols = parse_cols_with_header(self.key_cols_raw, header_map)
        self.calc_cols = parse_cols_with_header(self.calc_cols_raw, header_map)
        if self.dimension_cols_raw:
            self.dimension_cols = parse_cols_with_header(self.dimension_cols_raw, header_map)

        header_result = self.detector.detect_from_workbook(wb)
        data_start_excel_row = header_result.get_data_start_excel_row()

        sheet_names = wb.sheetnames
        if self.sheets == 'all':
            target_sheets = sheet_names
        else:
            indices = [int(x.strip()) for x in self.sheets.split(',') if x.strip().isdigit()]
            target_sheets = [sheet_names[i-1] for i in indices if 1 <= i <= len(sheet_names)]

        for sheet_name in target_sheets:
            ws = wb[sheet_name]
            self._process_sheet(ws, data_start_excel_row)

        out_name = generate_output_filename('subtotal', self.original_filename)
        out_dir = os.path.dirname(self.filepath)
        out_path = os.path.join(out_dir, out_name)
        wb.save(out_path)
        wb.close()
        return out_path

    def _process_sheet(self, ws, data_start_row):
        max_col = ws.max_column
        max_row = ws.max_row
        if data_start_row > max_row:
            return

        groups = {}
        for row_idx in range(data_start_row, max_row + 1):
            key = tuple(ws.cell(row_idx, col).value for col in self.key_cols)
            # ---- 修复：跳过主键全为空的行（避免生成无意义的小计行） ----
            if all(v is None or str(v).strip() == '' for v in key):
                continue
            groups.setdefault(key, []).append(row_idx)

        # 如果没有有效组，直接返回
        if not groups:
            return

        sorted_groups = sorted(groups.items(), key=lambda x: max(x[1]), reverse=True)

        for key, rows in sorted_groups:
            last_row = max(rows)
            ws.insert_rows(last_row + 1)
            subtotal_row = last_row + 1

            for col_idx in range(1, max_col + 1):
                cell = ws.cell(subtotal_row, col_idx)
                cell.fill = SUBTOTAL_FILL
                cell.font = SUBTOTAL_FONT
                cell.border = THIN_BORDER

            for i, col in enumerate(self.key_cols):
                ws.cell(subtotal_row, col).value = key[i]

            for col in self.dimension_cols:
                ws.cell(subtotal_row, col).value = '合计'

            for col in self.calc_cols:
                col_letter = get_column_letter(col)
                func = self._get_agg_func(col_letter)
                values = []
                for r in rows:
                    val = ws.cell(r, col).value
                    try:
                        val = float(val)
                    except (TypeError, ValueError):
                        val = 0
                    values.append(val)

                if func == 'sum':
                    result = sum(values)
                elif func == 'avg':
                    result = sum(values) / len(values) if values else 0
                elif func == 'min':
                    result = min(values) if values else 0
                elif func == 'max':
                    result = max(values) if values else 0
                elif func == 'count':
                    result = len(values)
                else:
                    result = sum(values)
                ws.cell(subtotal_row, col).value = result
