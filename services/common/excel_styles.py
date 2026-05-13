# services/common/excel_styles.py
"""公共 Excel 样式定义，避免硬编码"""
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

# 字体
TITLE_FONT = Font(name='微软雅黑', bold=True, size=14)
HEADER_FONT_BOLD = Font(name='微软雅黑', bold=True, size=10)
HEADER_FONT_RED = Font(name='微软雅黑', bold=True, size=10, color="FF0000")
NORMAL_FONT = Font(name='微软雅黑', size=10)
BOLD_FONT = Font(name='微软雅黑', bold=True, size=10)
NEGATIVE_FONT = Font(name='微软雅黑', size=10, color="FF0000")

# 填充
GRAY_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
GREEN_FILL = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
RED_FILL = PatternFill(start_color="F4B7BE", end_color="F4B7BE", fill_type="solid")
LIGHT_GRAY_FILL = PatternFill(start_color="BFBFBF", end_color="BFBFBF", fill_type="solid")

# 边框
THIN_BORDER = Border(
    left=Side('thin'), right=Side('thin'),
    top=Side('thin'), bottom=Side('thin')
)

# 对齐
CENTER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_ALIGN = Alignment(horizontal='left', vertical='center')
