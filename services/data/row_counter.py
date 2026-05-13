import os
import logging
from openpyxl import load_workbook

logger = logging.getLogger(__name__)


class RowCounter:
    """通用文件行数统计器（不含标题行）"""
    @staticmethod
    def count_excel(filepath: str) -> int:
        try:
            wb = load_workbook(filepath, read_only=True)
            rows = sum(ws.max_row - 1 for ws in wb.worksheets)
            wb.close()
            return max(rows, 0)
        except Exception as e:
            logger.warning(f"统计 Excel 行数失败 {filepath}: {e}")
            return 0

    @staticmethod
    def count_csv(filepath: str, encoding='utf-8') -> int:
        try:
            with open(filepath, 'r', encoding=encoding) as f:
                return sum(1 for _ in f) - 1
        except Exception as e:
            logger.warning(f"统计 CSV 行数失败 {filepath}: {e}")
            return 0

    @classmethod
    def count(cls, filepath: str) -> int:
        ext = os.path.splitext(filepath)[1].lower()
        if ext in ('.xlsx', '.xls', '.xlsm'):
            return cls.count_excel(filepath)
        return cls.count_csv(filepath)

