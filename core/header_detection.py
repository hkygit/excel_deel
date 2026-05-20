# core/header_detection.py
"""生产级Excel表头检测器"""
import re
import os
import time
from typing import List, Tuple, Optional, Dict, Any
from dataclasses import dataclass, field, asdict
from openpyxl import load_workbook
import logging
import hashlib

logger = logging.getLogger(__name__)

@dataclass
class HeaderDetectionConfig:
    max_check_rows: int = 20
    header_row_range: Tuple[int, int] = (0, 2)
    min_header_cols: int = 3
    max_header_cols: int = 200
    min_non_empty_ratio: float = 0.6
    min_string_ratio: float = 0.7
    max_duplicate_ratio: float = 0.3
    min_unique_cols: int = 2
    known_header_keywords: List[str] = field(default_factory=lambda: [
        "名称", "姓名", "编号", "代码", "日期", "时间", "金额", "数量",
        "单位", "备注", "状态", "类型", "级别", "部门", "岗位",
        "科目", "凭证", "摘要", "借方", "贷方", "余额", "本期", "上期",
        "病历号", "患者", "诊断", "药品", "剂量", "用法", "医生", "科室",
        "id", "name", "date", "time", "code", "type", "status", "amount"
    ])
    skip_rows_keywords: List[str] = field(default_factory=lambda: [
        "合计", "总计", "小计", "total", "sum", "汇总", "平均", "average"
    ])
    mandatory_columns: List[str] = field(default_factory=list)
    read_only: bool = True
    data_only: bool = True
    cache_ttl_seconds: int = 300

    def to_dict(self) -> Dict:
        return asdict(self)

    def get_hash(self) -> str:
        import json
        config_str = json.dumps(self.to_dict(), sort_keys=True)
        return hashlib.md5(config_str.encode()).hexdigest()[:8]

class HeaderDetectionResult:
    def __init__(self):
        self.header_rows: List[int] = []
        self.data_start_row: int = 0          # 0-based
        self.column_names: List[str] = []
        self.confidence: float = 0.0
        self.reason: str = ""
        self.metrics: Dict[str, Any] = {}
        self.warnings: List[str] = []
        self.timestamp: float = time.time()
        self.config_hash: str = ""

    def is_valid(self) -> bool:
        return bool(self.header_rows) and self.confidence >= 0.1

    def get_data_start_excel_row(self) -> int:
        """返回1-based的数据起始行"""
        return self.data_start_row + 1 if self.data_start_row >= 0 else 1

class SimpleExcelHeaderDetector:
    def __init__(self, config: HeaderDetectionConfig = None):
        self.config = config or HeaderDetectionConfig()
        self.config_hash = self.config.get_hash()
        self._cache = {}
        self._workbook_cache = {}
        self._keyword_set = {kw.lower() for kw in self.config.known_header_keywords}
        self._skip_set = {kw.lower() for kw in self.config.skip_rows_keywords}
        self._numeric_regex = re.compile(r'^[+-]?(\d{1,3}(,\d{3})*|\d+)(\.\d+)?%?$')
        self._invalid_chars_regex = re.compile(r'[\/*?\[\]]')

    def detect_headers(self, file_path: str, sheet_name: str = None, use_cache: bool = True) -> HeaderDetectionResult:
        self._validate_file(file_path)
        mtime = os.path.getmtime(file_path)
        cache_key = (file_path, mtime, self.config_hash, sheet_name)
        if use_cache and cache_key in self._cache:
            cached = self._cache[cache_key]
            if time.time() - cached.timestamp < self.config.cache_ttl_seconds:
                cached.metrics['cache_hit'] = True
                return cached

        result = HeaderDetectionResult()
        result.config_hash = self.config_hash
        try:
            wb = load_workbook(filename=file_path, read_only=self.config.read_only, data_only=self.config.data_only)
            ws = wb[sheet_name] if sheet_name else wb.active
            result = self._detect_from_worksheet(ws)
            wb.close()
        except Exception as e:
            logger.error(f"检测失败 {file_path}: {e}")
            result = self._create_fallback_result()
            result.warnings.append(str(e))

        if use_cache and result.is_valid():
            self._cache[cache_key] = result
        return result

    def detect_from_workbook(self, workbook, sheet_name: str = None) -> HeaderDetectionResult:
        workbook_id = id(workbook)
        cache_key = (workbook_id, sheet_name, self.config_hash)
        if cache_key in self._workbook_cache:
            cached = self._workbook_cache[cache_key]
            cached.metrics['workbook_cache_hit'] = True
            return cached

        try:
            ws = workbook[sheet_name] if sheet_name else workbook.active
            result = self._detect_from_worksheet(ws)
            result.config_hash = self.config_hash
        except Exception as e:
            logger.error(f"从workbook检测失败: {e}")
            result = self._create_fallback_result()
            result.warnings.append(str(e))

        self._workbook_cache[cache_key] = result
        return result

    def _validate_file(self, file_path):
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"文件不存在: {file_path}")
        if not os.path.isfile(file_path):
            raise ValueError(f"路径不是文件: {file_path}")
        ext = os.path.splitext(file_path)[1].lower()
        if ext not in {'.xlsx', '.xls', '.xlsm'}:
            raise ValueError(f"不支持的文件类型: {ext}")

    def _detect_from_worksheet(self, ws) -> HeaderDetectionResult:
        result = HeaderDetectionResult()
        start_time = time.time()
        max_row = min(ws.max_row, self.config.max_check_rows)
        max_col = min(ws.max_column, self.config.max_header_cols)
        if max_row == 0 or max_col == 0:
            result.warnings.append("工作表为空")
            return self._create_fallback_result()

        data = self._read_cell_range(ws, 1, max_row, 1, max_col)
        if not data:
            return self._create_fallback_result()

        candidates = []
        start, end = self.config.header_row_range
        for row_idx in range(min(end + 1, len(data))):
            if row_idx < start:
                continue
            candidate = self._evaluate_row(data, row_idx)
            if candidate['score'] > 0:
                candidates.append(candidate)

        if candidates:
            candidates.sort(key=lambda x: x['score'], reverse=True)
            best = candidates[0]
            result.header_rows = [best['row']]
            result.confidence = best['score']
            result.reason = best['reason']
            result.column_names = best['column_names']
            result.data_start_row = best['row'] + 1
        else:
            result = self._create_fallback_result()
            if data:
                result.column_names = [
                    str(cell).strip() if cell is not None and str(cell).strip()
                    else f"Col_{i+1}" for i, cell in enumerate(data[0])
                ]
        result.metrics.update({
            'execution_time_ms': (time.time() - start_time) * 1000,
            'rows_checked': len(data),
            'columns_checked': max_col,
            'candidates_found': len(candidates)
        })
        return result

    def _read_cell_range(self, ws, min_row, max_row, min_col, max_col):
        try:
            return [
                [cell for cell in row]
                for row in ws.iter_rows(min_row=min_row, max_row=max_row,
                                        min_col=min_col, max_col=max_col,
                                        values_only=True)
            ]
        except Exception as e:
            logger.warning(f"读取单元格失败: {e}")
            return []

    def _evaluate_row(self, data, row_idx):
        if row_idx >= len(data):
            return self._create_candidate(row_idx, 0, "索引越界")
        row_data = data[row_idx]
        if not row_data:
            return self._create_candidate(row_idx, 0, "空行")

        non_empty = [cell for cell in row_data if cell is not None and str(cell).strip() != '']
        non_empty_ratio = len(non_empty) / len(row_data) if row_data else 0
        if non_empty_ratio < 0.2:
            return self._create_candidate(row_idx, 0, f"非空{non_empty_ratio:.0%}")

        string_cells = []
        for cell in non_empty:
            s = str(cell).strip()
            if s and not self._is_numeric(s):
                string_cells.append(s)
        string_ratio = len(string_cells) / len(row_data) if row_data else 0
        unique_strings = set(string_cells)
        dup_ratio = 1 - len(unique_strings) / len(string_cells) if string_cells else 0

        keyword_matches = sum(1 for s in string_cells for kw in self._keyword_set if kw in s.lower())
        keyword_ratio = keyword_matches / len(string_cells) if string_cells else 0

        combined = ' '.join(string_cells).lower()
        if any(skip in combined for skip in self._skip_set):
            return self._create_candidate(row_idx, 0, "跳过关键词")

        valid_fmt = sum(1 for s in string_cells if self._is_valid_column_name(s))
        valid_fmt_ratio = valid_fmt / len(string_cells) if string_cells else 0

        score = 0.0
        reasons = []
        if non_empty_ratio >= self.config.min_non_empty_ratio:
            score += 0.3; reasons.append(f"非空{non_empty_ratio:.0%}")
        if string_ratio >= self.config.min_string_ratio:
            score += 0.3; reasons.append(f"文本{string_ratio:.0%}")
        if dup_ratio <= self.config.max_duplicate_ratio:
            score += 0.2; reasons.append(f"唯一{1-dup_ratio:.0%}")
        if keyword_ratio > 0.1:
            score += min(keyword_ratio, 0.2); reasons.append(f"关键词{keyword_ratio:.0%}")
        if valid_fmt_ratio > 0.5:
            score += 0.1; reasons.append("格式正确")
        score = max(0.0, min(1.0, score))

        column_names = [
            str(cell).strip() if cell is not None and str(cell).strip() != '' else f"Col_{i+1}"
            for i, cell in enumerate(row_data)
        ]
        return {
            'row': row_idx, 'score': score,
            'reason': ', '.join(reasons) or '基础特征',
            'column_names': column_names,
            'metrics': {}
        }

    def _create_candidate(self, row_idx, score, reason):
        return {'row': row_idx, 'score': score, 'reason': reason, 'column_names': [], 'metrics': {}}

    def _is_numeric(self, s: str) -> bool:
        s = s.strip().replace(',', '')
        if s.endswith('%'):
            s = s[:-1]
        try:
            float(s)
            return True
        except ValueError:
            return False

    def _is_valid_column_name(self, s: str) -> bool:
        if not s or len(s) > 50:
            return False
        if re.match(r'^[\d.,%\s-]+$', s):
            return False
        if self._invalid_chars_regex.search(s):
            return False
        return s.strip() != ''

    def _create_fallback_result(self) -> HeaderDetectionResult:
        result = HeaderDetectionResult()
        result.header_rows = [0]
        result.data_start_row = 1
        result.confidence = 0.3
        result.reason = "回退默认第一行"
        return result

    def clear_cache(self):
        self._cache.clear()
        self._workbook_cache.clear()


def detect_header_start(file_path: str, sheet_name: str = None) -> int:
    """返回数据起始行（1-based），失败返回2"""
    detector = SimpleExcelHeaderDetector()
    try:
        result = detector.detect_headers(file_path, sheet_name)
        if result.is_valid():
            return result.get_data_start_excel_row()
    except Exception:
        pass
    return 2
