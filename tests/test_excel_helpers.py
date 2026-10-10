# tests/test_excel_helpers.py
"""
针对 core.excel_helpers.normalize_excel_file 的回归测试。

P0 语义：
- 有效 .xlsx → 直接返回原路径（零复制）
- .xls / 损坏 .xlsx → 走转换路径，返回新的临时文件
- 尺寸超限 → 抛 ExcelNormalizationError
"""
import gc
import os
import time

import pytest
from openpyxl import Workbook, load_workbook

from core.excel_helpers import (
    normalize_excel_file,
    ExcelNormalizationError,
    MAX_EXCEL_ROWS,
    MAX_EXCEL_COLS,
)


def _make_simple_xlsx(path, n_rows: int = 2, n_cols: int = 2):
    """生成一个简单 Excel 文件，返回其路径。"""
    wb = Workbook()
    ws = wb.active
    ws.append([f"col{c + 1}" for c in range(n_cols)])
    for r in range(1, n_rows):
        ws.append([r * (c + 1) for c in range(n_cols)])
    wb.save(str(path))
    wb.close()
    return str(path)


def _robust_unlink(path: str, attempts: int = 5, delay: float = 0.1) -> None:
    """尽力清理系统临时目录下由 normalize_excel_file 生成的文件。

    Windows 上 openpyxl / zipfile 可能在 close() 后仍短暂持有文件句柄，
    因此做几次带 sleep 的重试；最后仍失败则吞掉异常。
    """
    gc.collect()
    for _ in range(attempts):
        if not os.path.exists(path):
            return
        try:
            os.unlink(path)
            return
        except PermissionError:
            gc.collect()
            time.sleep(delay)
    try:
        os.unlink(path)
    except PermissionError:
        pass


def _temp_normalized_files() -> set:
    """返回系统临时目录下所有 normalized_*.xlsx 文件名。"""
    import tempfile
    return {
        f for f in os.listdir(tempfile.gettempdir())
        if f.startswith("normalized_") and f.endswith(".xlsx")
    }


class TestNormalizeExcelFile:
    # ---------- P0 核心：有效 .xlsx 零复制 ----------
    def test_valid_xlsx_returns_original_path_and_is_intact(self, tmp_path):
        """P0 核心：有效 .xlsx 直接返回原路径，且内容不变。"""
        src = _make_simple_xlsx(tmp_path / "simple.xlsx", n_rows=5, n_cols=3)

        out = normalize_excel_file(src)

        # 关键断言：返回的应该就是原路径
        assert out == src, f"应返回原路径 {src}，实际 {out}"
        assert os.path.exists(out)

        # 内容应完整
        wb = load_workbook(out, read_only=True)
        try:
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            assert len(rows) == 5
            assert rows[0] == ('col1', 'col2', 'col3')
            assert rows[3] == (3, 6, 9)
        finally:
            wb.close()
        # 注意：out 就是 src，由 tmp_path fixture 自动清理，无需 unlink

    def test_valid_xlsx_does_not_create_temp_file(self, tmp_path):
        """P0 核心：有效 .xlsx 不创建任何 normalized_*.xlsx 临时文件。"""
        src = _make_simple_xlsx(tmp_path / "no_copy.xlsx", n_rows=3, n_cols=2)
        before = _temp_normalized_files()

        out = normalize_excel_file(src)

        after = _temp_normalized_files()
        new_files = after - before
        assert new_files == set(), f"不应创建临时文件: {new_files}"
        assert out == src

    def test_multiple_calls_no_temp_file_accumulation(self, tmp_path):
        """连续多次调用，每次都返回原路径，不累积临时文件。"""
        src = _make_simple_xlsx(tmp_path / "multi.xlsx", n_rows=3, n_cols=2)
        before = _temp_normalized_files()

        results = [normalize_excel_file(src) for _ in range(3)]

        after = _temp_normalized_files()
        new_files = after - before
        assert new_files == set(), f"多次调用不应产生临时文件: {new_files}"
        assert all(r == src for r in results)

    def test_below_limit_returns_original_path(self, tmp_path):
        """恰好等于上限的 .xlsx 仍应返回原路径。"""
        src = _make_simple_xlsx(tmp_path / "boundary.xlsx", n_rows=3, n_cols=2)

        out = normalize_excel_file(src, max_rows=3, max_cols=2)
        assert out == src

        wb = load_workbook(out, read_only=True)
        try:
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            assert len(rows) == 3
            assert len(rows[0]) == 2
        finally:
            wb.close()

    # ---------- 损坏文件仍走转换路径 ----------
    def test_corrupt_xlsx_does_not_return_original_path(self, tmp_path):
        """有 magic 但内部损坏的 .xlsx 不应直接返回原路径。"""
        bad = tmp_path / "corrupt.xlsx"
        # 有 ZIP magic，但内容是垃圾字节
        bad.write_bytes(b"PK\x03\x04" + b"\x00" * 200)

        out = normalize_excel_file(str(bad))

        # 关键断言：不能返回原路径（因为文件无效）
        assert out != str(bad), "损坏文件不应直接返回原路径"
        assert os.path.exists(out), "应返回一个可用的文件路径"
        _robust_unlink(out)

    # ---------- 基础异常 ----------
    def test_file_not_found(self):
        with pytest.raises(FileNotFoundError):
            normalize_excel_file("/nonexistent/path/to/file.xlsx")

    # ---------- 尺寸上限 ----------
    def test_exceed_rows_raises_instead_of_truncating(self, tmp_path):
        """行数超限时抛 ExcelNormalizationError。"""
        src = _make_simple_xlsx(tmp_path / "many_rows.xlsx", n_rows=11, n_cols=2)

        with pytest.raises(ExcelNormalizationError, match="行数"):
            normalize_excel_file(src, max_rows=5, max_cols=100)

    def test_exceed_cols_raises_instead_of_truncating(self, tmp_path):
        """列数超限时抛 ExcelNormalizationError。"""
        src = _make_simple_xlsx(tmp_path / "wide.xlsx", n_rows=2, n_cols=5)

        with pytest.raises(ExcelNormalizationError, match="列数"):
            normalize_excel_file(src, max_rows=100, max_cols=3)

    def test_no_temp_file_left_on_size_error(self, tmp_path):
        """尺寸超限时不应留下任何 normalized_*.xlsx 临时文件。"""
        src = _make_simple_xlsx(tmp_path / "big.xlsx", n_rows=20, n_cols=2)
        before = _temp_normalized_files()

        with pytest.raises(ExcelNormalizationError):
            normalize_excel_file(src, max_rows=5, max_cols=100)

        after = _temp_normalized_files()
        leaked = after - before
        assert leaked == set(), f"发现残留的临时文件: {leaked}"

    # ---------- 常量与异常类型 ----------
    def test_error_is_value_error_subclass(self):
        """向后兼容：ExcelNormalizationError 必须是 ValueError 的子类。"""
        assert issubclass(ExcelNormalizationError, ValueError)

    def test_default_limits_are_excel_hard_limits(self):
        """默认上限应是 Excel 硬限制，而不是旧版静默截断的 100000 / 100。"""
        assert MAX_EXCEL_ROWS == 1_048_576
        assert MAX_EXCEL_COLS == 16_384
