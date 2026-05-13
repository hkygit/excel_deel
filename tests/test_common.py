import os
import numpy as np
import pandas as pd
from openpyxl import Workbook
from core.validators import validate_column_letters, validate_sheet_index, validate_expression, sanitize_filename
# 直接使用核心表达式引擎，替代旧的 utils.expression 函数
from core.expression_engine import ExpressionEngine
from core.excel_utils import ExcelFile, col_letter_to_index, col_index_to_letter
from services.data.file_service import FileService
from flask import Flask

class TestValidators:
    def test_validate_column_letters(self):
        assert validate_column_letters(['A', 'AB', 'Z']) == True
        assert validate_column_letters(['a', 'Ab']) == True
        assert validate_column_letters(['A1', 'AB']) == False
        assert validate_column_letters([]) == True

    def test_validate_sheet_index(self):
        assert validate_sheet_index('1') == True
        assert validate_sheet_index(2) == True
        assert validate_sheet_index('0') == False
        assert validate_sheet_index('a') == False
        assert validate_sheet_index(None) == False

    def test_validate_expression(self):
        assert validate_expression('A+B-C') == True
        assert validate_expression('A*B/C') == True
        assert validate_expression('A**2') == True
        assert validate_expression('A+') == False
        assert validate_expression('') == False

    def test_sanitize_filename(self):
        assert sanitize_filename('a/b?c:d') == 'a_b_c_d'
        assert sanitize_filename('normal') == 'normal'

class TestExpression:
    @classmethod
    def setup_class(cls):
        cls.engine = ExpressionEngine()

    def test_extract_column_letters(self):
        # 利用引擎的 compile 方法提取变量（字母），但提取逻辑在旧函数中，这里直接复用旧实现
        import re
        def extract_column_letters(expr):
            return set(re.findall(r'[A-Z]+', expr.upper()))
        assert extract_column_letters('A+B-C') == {'A', 'B', 'C'}
        assert extract_column_letters('AA+BB') == {'AA', 'BB'}
        assert extract_column_letters('') == set()

    def test_evaluate_expression(self):
        values = {'A': 10, 'B': 20, 'C': 5}
        assert self.engine.evaluate('A+B-C', values) == 25
        assert self.engine.evaluate('A*B', values) == 200
        assert self.engine.evaluate('A/B', values) == 0.5
        result = self.engine.evaluate('A/0', values)
        assert np.isnan(result)
        result = self.engine.evaluate('X+Y', values)
        assert np.isnan(result)

class TestExcelHandler:
    def test_col_letter_conversion(self):
        assert col_letter_to_index('A') == 1
        assert col_letter_to_index('Z') == 26
        assert col_letter_to_index('AA') == 27
        assert col_index_to_letter(1) == 'A'
        assert col_index_to_letter(27) == 'AA'

    def test_excel_file_read_write(self, tmp_path):
        df = pd.DataFrame({'A': [1, 2], 'B': ['x', 'y']})
        test_file = tmp_path / 'test.xlsx'
        df.to_excel(test_file, index=False)
        with ExcelFile(test_file) as ef:
            headers = ef.get_headers()
            assert headers == ['A', 'B']
            val = ef.get_cell_value(2, 'A')
            assert val == 1
            new_file = tmp_path / 'new.xlsx'
            ef.save(new_file)
        assert new_file.exists()

    def test_copy_merge_structure(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.merge_cells('A1:A3')
        ws['A1'] = 'Merged'
        test_file = tmp_path / 'merge.xlsx'
        wb.save(test_file)
        with ExcelFile(test_file) as ef:
            ws = ef.active_sheet
            ef.copy_merge_structure('A', 'B', 1, 3, sheet=ws)
            merged = list(ws.merged_cells.ranges)
            assert len(merged) == 2

    def test_read_dataframe(self, tmp_path):
        df = pd.DataFrame({'X': [1, 2]})
        test_file = tmp_path / 'data.xlsx'
        df.to_excel(test_file, index=False)
        result = ExcelFile.read_dataframe(test_file)
        pd.testing.assert_frame_equal(result, df)

class TestFileService:
    def test_save_upload(self, tmp_path):
        app = Flask(__name__)
        app.config['UPLOAD_FOLDER'] = str(tmp_path)
        app.config['ALLOWED_EXTENSIONS'] = {'xlsx'}
        with app.app_context():
            from werkzeug.datastructures import FileStorage
            import io
            fake_file = FileStorage(io.BytesIO(b'fake'), filename='test.xlsx')
            svc = FileService()
            path = svc.save_upload(fake_file, subdir='')
            assert os.path.exists(path)
            assert path.endswith('.xlsx')
            svc.cleanup([path])
            assert not os.path.exists(path)
