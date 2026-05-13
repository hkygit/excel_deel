import io
from openpyxl import load_workbook

class TestCalculatorIntegration:
    def test_column_calculation(self, client, sample_excel_file):
        data = {
            'file': (sample_excel_file, 'test.xlsx'),
            'calc_type': 'column',
            'expression': 'A+B',
            'result_col_column': 'D',
            'merge_by_cols': ''
        }
        response = client.post('/calc/process', data=data, content_type='multipart/form-data')
        assert response.status_code == 200
        assert response.headers['Content-Disposition'].startswith('attachment')
        assert '.xlsx' in response.headers['Content-Disposition']
        wb = load_workbook(io.BytesIO(response.data))
        ws = wb.active
        assert ws['D2'].value == 30
        assert ws['D3'].value == 8

    def test_aggregate_backfill(self, client, sample_excel_file):
        data = {
            'file': (sample_excel_file, 'test.xlsx'),
            'calc_type': 'aggregate',
            'group_cols': 'A',
            'agg_cols': 'B',
            'agg_func': 'sum',
            'result_col_aggregate': 'E'
        }
        response = client.post('/calc/process', data=data, content_type='multipart/form-data')
        assert response.status_code == 200
        assert '.xlsx' in response.headers['Content-Disposition']
        wb = load_workbook(io.BytesIO(response.data))
        ws = wb.active
        assert ws['E2'].value == 20
        assert ws['E3'].value == 3

    def test_missing_file_error(self, client):
        response = client.post('/calc/process', data={'calc_type': 'column'})
        assert response.status_code == 400
        assert not response.get_json()['success']
        assert '未上传文件' in response.get_json()['error']

    def test_invalid_expression_error(self, client, sample_excel_file):
        data = {
            'file': (sample_excel_file, 'test.xlsx'),
            'calc_type': 'column',
            'expression': '',
            'result_col_column': 'D'
        }
        response = client.post('/calc/process', data=data, content_type='multipart/form-data')
        assert response.status_code == 400
        assert not response.get_json()['success']
        assert '表达式' in response.get_json()['error']
