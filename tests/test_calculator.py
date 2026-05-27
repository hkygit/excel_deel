import pandas as pd
from openpyxl import Workbook
from services.excel.calculator import ColumnCalculator
from services.excel.aggregator import Aggregator
from core.excel_utils import ExcelFile


class TestColumnCalculator:
    def test_calculate_column(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws['A1'] = 'A'
        ws['B1'] = 'B'
        ws['C1'] = 'Result'
        ws['A2'] = 10
        ws['B2'] = 20
        ws['A3'] = 5
        ws['B3'] = 3
        test_file = tmp_path / 'test.xlsx'
        wb.save(test_file)

        with ExcelFile(test_file) as ef:
            calculator = ColumnCalculator(ef)
            calculator.calculate_column('A+B', 'C', merge_by_cols=None)
            ef.save(test_file)

        result_df = pd.read_excel(test_file)
        assert result_df.loc[0, 'Result'] == 30
        assert result_df.loc[1, 'Result'] == 8

    def test_calculate_column_with_merge(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws['A1'] = 'Group'
        ws['B1'] = 'Value'
        ws['C1'] = 'Result'
        ws['A2'] = 'X'
        ws['B2'] = 10
        ws['A3'] = 'X'
        ws['B3'] = 20
        ws['A4'] = 'Y'
        ws['B4'] = 5
        test_file = tmp_path / 'test.xlsx'
        wb.save(test_file)

        with ExcelFile(test_file) as ef:
            calculator = ColumnCalculator(ef)
            calculator.calculate_column('B*2', 'C', merge_by_cols=['A'])
            ef.save(test_file)

        result_df = pd.read_excel(test_file)
        assert result_df.loc[0, 'Result'] == 20
        assert pd.isna(result_df.loc[1, 'Result'])
        assert result_df.loc[2, 'Result'] == 10


class TestAggregator:
    def test_aggregate_sum(self):
        df = pd.DataFrame({'Group': ['X', 'X', 'Y'], 'Value': [10, 20, 5]})
        result = Aggregator.simple_aggregate(df, ['Group'], ['Value'], 'sum')
        assert len(result) == 2
        assert result.loc[result['Group'] == 'X', 'Value'].iloc[0] == 30
        assert result.loc[result['Group'] == 'Y', 'Value'].iloc[0] == 5

    def test_aggregate_mean(self):
        df = pd.DataFrame({'Group': ['X', 'X'], 'Value': [10, 30]})
        result = Aggregator.simple_aggregate(df, ['Group'], ['Value'], 'mean')
        assert result.loc[0, 'Value'] == 20

    def test_aggregate_by_keys(self):
        df = pd.DataFrame({
            'key': ['A', 'A', 'B'],
            'sum_col': [10, 20, 5],
            'other_col': ['first', 'second', 'other']
        })
        result = Aggregator.aggregate_by_keys(df, ['key'], ['sum_col'], ['other_col'])
        assert len(result) == 2
        assert result.loc[result['key'] == 'A', 'sum_col'].iloc[0] == 30
        assert result.loc[result['key'] == 'A', 'other_col'].iloc[0] == 'first'
        assert result.loc[result['key'] == 'B', 'sum_col'].iloc[0] == 5
