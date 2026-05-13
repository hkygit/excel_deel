import pytest
import tempfile
import os
from openpyxl import Workbook, load_workbook
from blueprints.splitter import split_excel_to_temp, sanitize_filename, filter_workbook_inplace
import pandas as pd

def create_test_excel(filepath):
    wb = Workbook()
    ws = wb.active
    ws['A1'] = 'Category'
    ws['B1'] = 'Value'
    ws['A2'] = 'A'
    ws['B2'] = 10
    ws['A3'] = 'A'
    ws['B3'] = 20
    ws['A4'] = 'B'
    ws['B4'] = 5
    wb.save(filepath)
    return filepath

class TestSplitter:
    def test_split_excel_to_temp(self, tmp_path):
        input_file = tmp_path / 'input.xlsx'
        create_test_excel(input_file)
        output_dir = tmp_path / 'output'
        output_dir.mkdir()
        files = split_excel_to_temp(str(input_file), split_column_name='Category', include_values=[], exclude_values=[], skip_sheet_indices=set(), temp_dir=str(output_dir))
        assert len(files) == 2
        assert any('A.xlsx' in f for f in files)
        assert any('B.xlsx' in f for f in files)

    def test_sanitize_filename(self):
        assert sanitize_filename('a/b?c') == 'a_b_c'
        assert sanitize_filename('normal') == 'normal'

    def test_filter_with_skip_sheets(self, tmp_path):
        """测试原生筛选 + 跳过索引，使用 openpyxl 验证结果"""
        wb = Workbook()
        ws1 = wb.active
        ws1.title = 'Main'
        ws1['A1'] = 'Name'
        ws1['B1'] = 'Value'
        ws1['A2'] = 'Alice'
        ws1['B2'] = '100'
        ws2 = wb.create_sheet('Second')
        ws2['A1'] = 'Other'
        ws2['B1'] = 'Value'
        ws2['A2'] = 'should_not_filter'
        in_path = tmp_path / 'test.xlsx'
        wb.save(in_path)

        form = {
            'condition_field_0': 'Name',
            'condition_op_0': '=',
            'condition_value_0': 'Alice'
        }
        # 调用新函数，传入跳过索引 {2}
        filtered_path = filter_workbook_inplace(str(in_path), form, skip_sheet_indices={2})

        # 加载筛选后的工作簿
        wb_result = load_workbook(filtered_path)

        # 第一个 sheet 应该只剩下表头 + Alice 行
        ws1_res = wb_result['Main']
        assert ws1_res.max_row == 2
        assert ws1_res.cell(row=2, column=1).value == 'Alice'
        assert ws1_res.cell(row=2, column=2).value == '100'

        # 第二个 sheet 应被跳过，原始数据保留
        ws2_res = wb_result['Second']
        assert ws2_res.max_row == 2
        assert ws2_res.cell(row=2, column=1).value == 'should_not_filter'
        wb_result.close()
