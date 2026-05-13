import pytest
import tempfile
import os
import shutil
from openpyxl import Workbook
import pandas as pd
from app import create_app


@pytest.fixture
def app():
    """每个测试函数独立创建一个 Flask 应用和临时目录，并在结束后清理。"""
    app = create_app()
    upload = tempfile.mkdtemp()
    simplify_upload = tempfile.mkdtemp()
    simplify_output = tempfile.mkdtemp()
    simplify_status = tempfile.mkdtemp()
    app.config.update({
        "TESTING": True,
        "UPLOAD_FOLDER": upload,
        "SIMPLIFY_UPLOAD_FOLDER": simplify_upload,
        "SIMPLIFY_OUTPUT_FOLDER": simplify_output,
        "SIMPLIFY_TASK_STATUS_FOLDER": simplify_status,
        "WTF_CSRF_ENABLED": False,
    })
    yield app
    # 清理所有临时目录
    for d in [upload, simplify_upload, simplify_output, simplify_status]:
        shutil.rmtree(d, ignore_errors=True)


@pytest.fixture
def client(app):
    return app.test_client()


@pytest.fixture
def sample_excel_file():
    from io import BytesIO
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "A"
    ws["B1"] = "B"
    ws["C1"] = "C"
    ws["A2"] = 10
    ws["B2"] = 20
    ws["C2"] = 30
    ws["A3"] = 5
    ws["B3"] = 3
    ws["C3"] = 2
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


@pytest.fixture
def sample_csv_header_file():
    from io import BytesIO
    content = "key,value,other\n"
    return BytesIO(content.encode("utf-8"))


@pytest.fixture
def sample_csv_data_file():
    from io import BytesIO
    content = "A,10,x\nA,20,y\nB,5,z\n"
    return BytesIO(content.encode("utf-8"))


@pytest.fixture
def sample_excel_comparison_files():
    from io import BytesIO
    df1 = pd.DataFrame({"id": [1, 2], "name": ["Alice", "Bob"], "score": [100, 90]})
    df2 = pd.DataFrame({"id": [1, 3], "name": ["Alice", "Charlie"], "score": [100, 80]})
    f1 = BytesIO()
    f2 = BytesIO()
    with pd.ExcelWriter(f1, engine="openpyxl") as w: df1.to_excel(w, index=False)
    with pd.ExcelWriter(f2, engine="openpyxl") as w: df2.to_excel(w, index=False)
    f1.seek(0); f2.seek(0)
    return f1, f2
