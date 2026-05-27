import io
import pytest
import pandas as pd
from openpyxl import Workbook

class TestMatcherIntegration:
    @pytest.fixture
    def target_file(self):
        wb = Workbook()
        ws = wb.active
        ws['A1'] = 'ID'
        ws['B1'] = 'Value1'
        ws['C1'] = 'Value2'
        ws['A2'] = 1
        ws['B2'] = ''
        ws['C2'] = ''
        ws['A3'] = 2
        ws['B3'] = ''
        ws['C3'] = ''
        ws['A4'] = 3
        ws['B4'] = ''
        ws['C4'] = ''
        f = io.BytesIO()
        wb.save(f)
        f.seek(0)
        return f

    @pytest.fixture
    def target_file_duplicate(self):
        wb = Workbook()
        ws = wb.active
        ws['A1'] = 'ID'
        ws['B1'] = 'Value1'
        ws['C1'] = 'Value2'
        ws['A2'] = 1
        ws['B2'] = ''
        ws['C2'] = ''
        ws['A3'] = 1
        ws['B3'] = ''
        ws['C3'] = ''
        ws['A4'] = 2
        ws['B4'] = ''
        ws['C4'] = ''
        f = io.BytesIO()
        wb.save(f)
        f.seek(0)
        return f

    @pytest.fixture
    def source_file(self):
        df = pd.DataFrame({'ID': [1, 2, 3], 'Score1': [100, 200, 300], 'Score2': [101, 202, 303]})
        f = io.BytesIO()
        with pd.ExcelWriter(f, engine='openpyxl') as w:
            df.to_excel(w, index=False)
        f.seek(0)
        return f

    def test_match_single_column(self, client, target_file, source_file):
        """测试单列匹配（向后兼容）"""
        data = {
            'target_file': (target_file, 'target.xlsx'),
            'source_file': (source_file, 'source.xlsx'),
            'target_key_cols': 'A',
            'source_match_cols': 'A',
            'source_value_col': 'B',
            'result_col': 'B',
        }
        response = client.post('/match/process', data=data, content_type='multipart/form-data')
        assert response.status_code == 200
        result_df = pd.read_excel(io.BytesIO(response.data))
        assert result_df.loc[0, 'Value1'] == 100
        assert result_df.loc[1, 'Value1'] == 200
        assert result_df.loc[2, 'Value1'] == 300

    def test_match_multiple_columns(self, client, target_file, source_file):
        """测试多列匹配：源值列 B,C -> 结果列 B,C"""
        data = {
            'target_file': (target_file, 'target.xlsx'),
            'source_file': (source_file, 'source.xlsx'),
            'target_key_cols': 'A',
            'source_match_cols': 'A',
            'source_value_col': 'B,C',
            'result_col': 'B,C',
        }
        response = client.post('/match/process', data=data, content_type='multipart/form-data')
        assert response.status_code == 200
        result_df = pd.read_excel(io.BytesIO(response.data))
        assert result_df.loc[0, 'Value1'] == 100
        assert result_df.loc[1, 'Value1'] == 200
        assert result_df.loc[2, 'Value1'] == 300
        assert result_df.loc[0, 'Value2'] == 101
        assert result_df.loc[1, 'Value2'] == 202
        assert result_df.loc[2, 'Value2'] == 303

    def test_match_multiple_columns_with_new_result_names(self, client, target_file, source_file):
        """测试多列匹配，结果列为新建列名"""
        data = {
            'target_file': (target_file, 'target.xlsx'),
            'source_file': (source_file, 'source.xlsx'),
            'target_key_cols': 'A',
            'source_match_cols': 'A',
            'source_value_col': 'B,C',
            'result_col': '新分数1,新分数2',
        }
        response = client.post('/match/process', data=data, content_type='multipart/form-data')
        assert response.status_code == 200
        result_df = pd.read_excel(io.BytesIO(response.data))
        assert '新分数1' in result_df.columns
        assert '新分数2' in result_df.columns
        assert result_df.loc[0, '新分数1'] == 100
        assert result_df.loc[0, '新分数2'] == 101

    def test_match_column_count_mismatch(self, client, target_file, source_file):
        """测试源值列与结果列数量不一致时返回错误"""
        data = {
            'target_file': (target_file, 'target.xlsx'),
            'source_file': (source_file, 'source.xlsx'),
            'target_key_cols': 'A',
            'source_match_cols': 'A',
            'source_value_col': 'B,C',
            'result_col': 'B',   # 数量不匹配
        }
        response = client.post('/match/process', data=data, content_type='multipart/form-data')
        assert response.status_code == 400
        resp_json = response.get_json()
        assert not resp_json['success']
        assert '数量' in resp_json['error']

    def test_match_with_duplicate_keys_disabled(self, client, target_file_duplicate, source_file):
        data = {
            'target_file': (target_file_duplicate, 'target_dup.xlsx'),
            'source_file': (source_file, 'source.xlsx'),
            'target_key_cols': 'A',
            'source_match_cols': 'A',
            'source_value_col': 'B,C',
            'result_col': 'B,C',
        }
        response = client.post('/match/process', data=data, content_type='multipart/form-data')
        assert response.status_code == 400
        resp_json = response.get_json()
        assert not resp_json['success']
        assert '重复' in resp_json['error']

    def test_match_with_duplicate_keys_enabled(self, client, target_file_duplicate, source_file):
        data = {
            'target_file': (target_file_duplicate, 'target_dup.xlsx'),
            'source_file': (source_file, 'source.xlsx'),
            'target_key_cols': 'A',
            'source_match_cols': 'A',
            'source_value_col': 'B,C',
            'result_col': 'B,C',
            'allow_duplicate_keys': 'on',
        }
        response = client.post('/match/process', data=data, content_type='multipart/form-data')
        assert response.status_code == 200
        result_df = pd.read_excel(io.BytesIO(response.data))
        # 第一行填充，第二行（重复键）留空
        assert result_df.loc[0, 'Value1'] == 100
        assert result_df.loc[1, 'Value1'] == '' or pd.isna(result_df.loc[1, 'Value1'])
        assert result_df.loc[2, 'Value1'] == 200

    def test_debug_mode_multiple_columns(self, client, target_file, source_file):
        data = {
            'target_file': (target_file, 'target.xlsx'),
            'source_file': (source_file, 'source.xlsx'),
            'target_key_cols': 'A',
            'source_match_cols': 'A',
            'source_value_col': 'B,C',
            'result_col': 'B,C',
        }
        response = client.post('/match/debug', data=data, content_type='multipart/form-data')
        assert response.status_code == 200
        content = response.data.decode('utf-8')
        assert 'source_value_cols' in content or 'config' in content
