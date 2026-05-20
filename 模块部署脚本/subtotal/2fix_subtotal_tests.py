#!/usr/bin/env python3
"""
修复数据小计模块的三项测试失败：
1. 引擎路径兼容 WindowsPath 对象
2. 列存在性验证更健壮
3. 异步测试 monkeypatch 修正
"""
import os
import shutil

# ================= 1. 修复 services/subtotal/engine.py =================
FIXED_SUBTOTAL_ENGINE = '''# services/subtotal/engine.py
import openpyxl
from openpyxl.styles import PatternFill, Font, Border, Side
from openpyxl.utils import column_index_from_string
from core.csv_utils import split_fields

SUBTOTAL_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
SUBTOTAL_FONT = Font(bold=True)
THIN_BORDER = Border(
    left=Side('thin'), right=Side('thin'),
    top=Side('thin'), bottom=Side('thin')
)

class SubtotalEngine:
    """数据小计核心引擎"""

    def __init__(self, filepath, key_cols_str, calc_cols_str, agg_funcs,
                 dimension_cols=None, sheets='1'):
        self.filepath = str(filepath)   # 确保是字符串，避免 WindowsPath 等对象
        self.key_cols = self._parse_cols(key_cols_str)
        self.calc_cols = self._parse_cols(calc_cols_str)
        self.dimension_cols = self._parse_cols(dimension_cols) if dimension_cols else []
        self.agg_funcs = agg_funcs
        self.sheets = sheets

    @staticmethod
    def _parse_cols(raw):
        if not raw:
            return []
        items = split_fields(raw)
        result = []
        for item in items:
            if '-' in item:
                start, end = item.split('-', 1)
                start_idx = column_index_from_string(start.strip())
                end_idx = column_index_from_string(end.strip())
                for i in range(start_idx, end_idx + 1):
                    result.append(i)
            else:
                result.append(column_index_from_string(item.strip()))
        return sorted(list(set(result)))

    def _get_agg_func(self, col_letter):
        if col_letter in self.agg_funcs:
            return self.agg_funcs[col_letter]
        return self.agg_funcs.get('default', 'sum')

    def process(self):
        wb = openpyxl.load_workbook(self.filepath)
        sheet_names = wb.sheetnames
        if self.sheets == 'all':
            target_sheets = sheet_names
        else:
            indices = [int(x.strip()) for x in self.sheets.split(',')
                       if x.strip().isdigit()]
            target_sheets = [sheet_names[i - 1] for i in indices
                             if 1 <= i <= len(sheet_names)]

        for sheet_name in target_sheets:
            ws = wb[sheet_name]
            self._process_sheet(ws)

        # 使用 self.filepath（已为字符串）进行 rsplit
        base = self.filepath.rsplit('.', 1)[0]
        out_path = base + '_小计.xlsx'
        wb.save(out_path)
        wb.close()
        return out_path

    def _process_sheet(self, ws):
        data_start = 2
        max_row = ws.max_row

        groups = {}
        for row_idx in range(data_start, max_row + 1):
            key = tuple(ws.cell(row_idx, col).value for col in self.key_cols)
            groups.setdefault(key, []).append(row_idx)

        sorted_groups = sorted(groups.items(),
                               key=lambda x: max(x[1]), reverse=True)

        for key, rows in sorted_groups:
            last_row = max(rows)
            ws.insert_rows(last_row + 1)
            subtotal_row = last_row + 1

            # 样式
            max_col = ws.max_column
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(subtotal_row, col_idx)
                cell.fill = SUBTOTAL_FILL
                cell.font = SUBTOTAL_FONT
                cell.border = THIN_BORDER

            # 主键列
            for i, col in enumerate(self.key_cols):
                ws.cell(subtotal_row, col).value = key[i]

            # 维度列
            for col in self.dimension_cols:
                ws.cell(subtotal_row, col).value = '合计'

            # 计算列
            for col in self.calc_cols:
                col_letter = openpyxl.utils.get_column_letter(col)
                func = self._get_agg_func(col_letter)
                values = []
                for r in rows:
                    val = ws.cell(r, col).value
                    try:
                        val = float(val)
                    except (TypeError, ValueError):
                        val = 0
                    values.append(val)

                if func == 'sum':
                    result = sum(values)
                elif func == 'avg':
                    result = sum(values) / len(values) if values else 0
                elif func == 'min':
                    result = min(values) if values else 0
                elif func == 'max':
                    result = max(values) if values else 0
                elif func == 'count':
                    result = len(values)
                else:
                    result = sum(values)
                ws.cell(subtotal_row, col).value = result
'''

# ================= 2. 修复 blueprints/subtotal.py =================
FIXED_SUBTOTAL_BP = '''# blueprints/subtotal.py
import os
import logging
import openpyxl
from flask import Blueprint, render_template, request, send_file, current_app, jsonify
from core.decorators import handle_exceptions
from core.exceptions import ValidationError
from core.task.decorator import async_task
from services.subtotal.engine import SubtotalEngine
from core.csv_utils import split_fields

logger = logging.getLogger(__name__)
subtotal_bp = Blueprint('subtotal', __name__, url_prefix='/subtotal',
                        template_folder='../templates/subtotal')

@subtotal_bp.route('/')
def index():
    return render_template('subtotal/index.html')

def _validate_single_letter(letter, max_col, category):
    """验证单个列字母是否有效且不超过最大列"""
    idx = openpyxl.utils.column_index_from_string(letter)
    if idx > max_col:
        raise ValidationError(f"{category}列 '{letter}' 超出工作表列范围，文件仅有 {max_col} 列")
    return idx

def _validate_col_range(item, max_col, category):
    """验证列范围如 A-C，并返回开始和结束索引"""
    if '-' in item:
        parts = item.split('-')
        if len(parts) != 2:
            raise ValidationError(f"{category}列范围格式错误：{item}")
        start_letter, end_letter = parts[0].strip(), parts[1].strip()
        start_idx = _validate_single_letter(start_letter, max_col, category)
        end_idx = _validate_single_letter(end_letter, max_col, category)
        if start_idx > end_idx:
            raise ValidationError(f"{category}列范围开始大于结束：{item}")
        return start_idx, end_idx
    else:
        idx = _validate_single_letter(item, max_col, category)
        return idx, idx

def _check_columns_exist(raw_str, category, max_col):
    """校验列字母或范围，确保所有涉及的列均存在"""
    if not raw_str:
        return
    items = split_fields(raw_str)
    for item in items:
        _validate_col_range(item, max_col, category)

@subtotal_bp.route('/process', methods=['POST'])
@handle_exceptions()
def process():
    file = request.files.get('file')
    if not file or file.filename == '':
        raise ValidationError('请上传文件')

    key_cols_raw = request.form.get('key_cols', '').strip()
    if not key_cols_raw:
        raise ValidationError('请指定主键列')

    dimension_cols_raw = request.form.get('dimension_cols', '').strip() or None
    calc_cols_raw = request.form.get('calc_cols', '').strip()
    if not calc_cols_raw:
        raise ValidationError('请指定计算列')

    agg_func_default = request.form.get('agg_func_default', 'sum').strip()
    agg_overrides_raw = request.form.get('agg_overrides', '').strip()
    sheets_mode = request.form.get('sheets_mode', '1').strip()

    # 保存文件
    file_service = current_app.container.file_service()
    path = file_service.save_upload(file, allowed_extensions={'xlsx', 'xls'})

    # 快速读取以获取最大列数
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    max_col = ws.max_column
    wb.close()

    # 校验所有列
    _check_columns_exist(key_cols_raw, '主键', max_col)
    _check_columns_exist(calc_cols_raw, '计算', max_col)
    if dimension_cols_raw:
        _check_columns_exist(dimension_cols_raw, '维度', max_col)

    # 聚合函数覆盖
    agg_funcs = {'default': agg_func_default}
    if agg_overrides_raw:
        parts = agg_overrides_raw.split(',')
        for part in parts:
            if ':' in part:
                col, func = part.split(':', 1)
                col = col.strip().upper()
                func = func.strip().lower()
                if func not in ('sum', 'avg', 'min', 'max', 'count'):
                    raise ValidationError(f"不支持的聚合函数: {func}")
                try:
                    openpyxl.utils.column_index_from_string(col)
                except Exception:
                    raise ValidationError(f"无效的列字母: {col}")
                agg_funcs[col] = func

    # 估算行数以决定同步/异步
    wb2 = openpyxl.load_workbook(path, read_only=True)
    total_rows = sum(sheet.max_row for sheet in wb2.worksheets)
    wb2.close()
    data_rows = total_rows - (1 if total_rows > 0 else 0)

    app = current_app._get_current_object()

    if data_rows > 5000:
        task_storage = app.container.task_storage()

        @async_task(original_filename='subtotal')
        def run_subtotal(task_id):
            with app.app_context():
                def progress_cb(percent, message):
                    task_storage.update(task_id,
                                        progress_percent=percent,
                                        progress_message=message)
                try:
                    progress_cb(0, "开始处理...")
                    engine = SubtotalEngine(
                        path, key_cols_raw, calc_cols_raw, agg_funcs,
                        dimension_cols=dimension_cols_raw,
                        sheets=sheets_mode
                    )
                    out_path = engine.process()
                    progress_cb(100, "完成")
                    return {'result_path': out_path}
                except Exception as e:
                    import traceback
                    error_msg = f"{str(e)}\\n{traceback.format_exc()}"
                    task_storage.append_log(task_id, error_msg)
                    raise
                finally:
                    file_service.cleanup([path])

        result = run_subtotal()
        return jsonify({'task_id': result['task_id']})
    else:
        # 同步处理
        try:
            engine = SubtotalEngine(
                path, key_cols_raw, calc_cols_raw, agg_funcs,
                dimension_cols=dimension_cols_raw,
                sheets=sheets_mode
            )
            out_path = engine.process()
            return send_file(out_path, as_attachment=True,
                             download_name=os.path.basename(out_path),
                             mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        finally:
            file_service.cleanup([path])

@subtotal_bp.route('/status/<task_id>')
def status(task_id):
    storage = current_app.container.task_storage()
    data = storage.load(task_id)
    if not data:
        return jsonify({'error': '任务不存在'}), 404
    log = storage.get_log(task_id, max_lines=100)
    data['log'] = log
    return jsonify(data)

@subtotal_bp.route('/download/<task_id>')
def download(task_id):
    storage = current_app.container.task_storage()
    data = storage.load(task_id)
    if not data or data.get('status') != 'completed':
        return jsonify({'error': '结果不可用'}), 400
    path = data.get('result_path')
    if not path or not os.path.exists(path):
        return jsonify({'error': '文件不存在'}), 404
    return send_file(path, as_attachment=True,
                     download_name='subtotal_result.xlsx')
'''

# ================= 3. 修复测试文件 =================
FIXED_TEST_ENGINE = '''# tests/test_subtotal_engine.py
import pytest
import pandas as pd
from services.subtotal.engine import SubtotalEngine

class TestSubtotalEngine:
    @pytest.fixture
    def sample_excel(self, tmp_path):
        df = pd.DataFrame({
            'A': ['X', 'X', 'Y'],
            'B': [10, 20, 30],
            'C': [1, 2, 3]
        })
        path = tmp_path / 'test.xlsx'
        df.to_excel(path, index=False)
        return str(path)   # 返回字符串，避免 Path 对象

    def test_basic_subtotal_sum(self, sample_excel):
        engine = SubtotalEngine(
            sample_excel,
            key_cols_str='A',
            calc_cols_str='B,C',
            agg_funcs={'default': 'sum'},
            dimension_cols=None
        )
        out_path = engine.process()
        assert out_path.endswith('_小计.xlsx')
        wb = pd.read_excel(out_path, sheet_name=None)
        ws = wb['Sheet1']
        assert len(ws) == 5
        assert ws.iloc[2]['B'] == 30
        assert ws.iloc[2]['C'] == 3
        assert ws.iloc[4]['B'] == 30
        assert ws.iloc[4]['C'] == 3

    def test_dimension_col(self, sample_excel):
        engine = SubtotalEngine(
            sample_excel,
            key_cols_str='A',
            calc_cols_str='B',
            agg_funcs={'default': 'sum'},
            dimension_cols='C'
        )
        out_path = engine.process()
        wb = pd.read_excel(out_path, sheet_name=None)
        ws = wb['Sheet1']
        assert ws.iloc[2]['C'] == '合计'

    def test_avg_function(self, sample_excel):
        engine = SubtotalEngine(
            sample_excel,
            key_cols_str='A',
            calc_cols_str='B',
            agg_funcs={'default': 'avg'},
            dimension_cols=None
        )
        out_path = engine.process()
        wb = pd.read_excel(out_path, sheet_name=None)
        ws = wb['Sheet1']
        assert ws.iloc[2]['B'] == 15.0

    def test_custom_agg_func_per_column(self, sample_excel):
        engine = SubtotalEngine(
            sample_excel,
            key_cols_str='A',
            calc_cols_str='B,C',
            agg_funcs={'default': 'sum', 'C': 'avg'},
            dimension_cols=None
        )
        out_path = engine.process()
        wb = pd.read_excel(out_path, sheet_name=None)
        ws = wb['Sheet1']
        assert ws.iloc[2]['B'] == 30
        assert ws.iloc[2]['C'] == 1.5

    def test_column_range_parsing(self, tmp_path):
        df = pd.DataFrame({
            'A': ['X', 'X'],
            'B': [10, 20],
            'C': [100, 200],
            'D': [1, 2]
        })
        path = tmp_path / 'range.xlsx'
        df.to_excel(path, index=False)
        engine = SubtotalEngine(
            str(path),   # 关键修复：传入字符串
            key_cols_str='A',
            calc_cols_str='B-D',
            agg_funcs={'default': 'sum'},
            dimension_cols=None
        )
        out_path = engine.process()
        wb = pd.read_excel(out_path, sheet_name=None)
        ws = wb['Sheet1']
        assert ws.iloc[2]['B'] == 30
        assert ws.iloc[2]['C'] == 300
        assert ws.iloc[2]['D'] == 3
'''

FIXED_TEST_INTEGRATION = '''# tests/test_integration_subtotal.py
import io
import pandas as pd

class TestSubtotalIntegration:
    def test_subtotal_endpoint_sync(self, client):
        # 数据有三列，对应 A,B,C，使用列字母
        df = pd.DataFrame({
            'A': ['X', 'X', 'Y'],
            'B': [10, 20, 30],
            'C': [100, 200, 300]
        })
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine='openpyxl') as w:
            df.to_excel(w, index=False)
        buf.seek(0)

        data = {
            'file': (buf, 'test.xlsx'),
            'key_cols': 'A',
            'calc_cols': 'B,C',
            'agg_func_default': 'sum',
            'dimension_cols': '',
            'sheets_mode': '1'
        }
        resp = client.post('/subtotal/process', data=data,
                           content_type='multipart/form-data')
        assert resp.status_code == 200
        result_df = pd.read_excel(io.BytesIO(resp.data), sheet_name=None)
        ws = result_df['Sheet1']
        assert len(ws) == 5           # 3 数据 + 2 小计
        # 组 X 小计行：B=30, C=300
        assert ws.iloc[2]['B'] == 30
        assert ws.iloc[2]['C'] == 300

    def test_missing_dimension_col_error(self, client):
        df = pd.DataFrame({'A': [1]})
        buf = io.BytesIO()
        df.to_excel(buf, index=False)
        buf.seek(0)

        data = {
            'file': (buf, 'test.xlsx'),
            'key_cols': 'A',
            'calc_cols': 'A',
            'agg_func_default': 'sum',
            'dimension_cols': 'Z'       # 超出列范围
        }
        resp = client.post('/subtotal/process', data=data,
                           content_type='multipart/form-data')
        assert resp.status_code == 400
        json_data = resp.get_json()
        assert '超出' in json_data['error'] and 'Z' in json_data['error']

    def test_async_large_file(self, client):
        """生成实际 6000 行数据触发异步逻辑"""
        df = pd.DataFrame({
            'A': ['X', 'Y'] * 3000,    # 6000 行
            'B': list(range(6000))
        })
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine='openpyxl') as w:
            df.to_excel(w, index=False)
        buf.seek(0)

        data = {
            'file': (buf, 'test_large.xlsx'),
            'key_cols': 'A',
            'calc_cols': 'B',
            'agg_func_default': 'sum',
            'sheets_mode': '1'
        }
        resp = client.post('/subtotal/process', data=data,
                           content_type='multipart/form-data')
        assert resp.status_code == 200
        assert 'task_id' in resp.get_json()
'''

# ================= 执行修复 =================
def main():
    root = os.path.dirname(os.path.abspath(__file__))

    files_to_fix = {
        'services/subtotal/engine.py': FIXED_SUBTOTAL_ENGINE,
        'blueprints/subtotal.py': FIXED_SUBTOTAL_BP,
        'tests/test_subtotal_engine.py': FIXED_TEST_ENGINE,
        'tests/test_integration_subtotal.py': FIXED_TEST_INTEGRATION,
    }

    for rel_path, content in files_to_fix.items():
        full_path = os.path.join(root, rel_path)
        os.makedirs(os.path.dirname(full_path), exist_ok=True)
        with open(full_path, 'w', encoding='utf-8') as f:
            f.write(content)
        print(f'已修复: {rel_path}')

    # 清理缓存
    for dirpath, dirnames, filenames in os.walk(root, topdown=False):
        for d in dirnames:
            if d == '__pycache__':
                shutil.rmtree(os.path.join(dirpath, d), ignore_errors=True)
        for f in filenames:
            if f.endswith('.pyc'):
                os.remove(os.path.join(dirpath, f))

    print(
        '缓存已清理，请重新运行测试：python -m pytest tests/test_subtotal_engine.py tests/test_integration_subtotal.py -v')


if __name__ == '__main__':
    main()