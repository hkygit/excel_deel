# blueprints/subtotal.py
import os
import logging
import openpyxl
from flask import Blueprint, render_template, request, send_file, current_app, jsonify
from core.decorators import handle_exceptions
from core.exceptions import ValidationError
from core.task.decorator import async_task
from services.subtotal.engine import SubtotalEngine, _build_header_map, resolve_single_col
from core.csv_utils import split_fields
from core.file_naming import generate_output_filename, build_content_disposition

logger = logging.getLogger(__name__)
subtotal_bp = Blueprint('subtotal', __name__, url_prefix='/subtotal',
                        template_folder='../templates/subtotal')

@subtotal_bp.route('/')
def index():
    return render_template('subtotal/index.html')

def _check_columns(raw, category, header_map, max_col):
    if not raw:
        return
    items = split_fields(raw)
    for item in items:
        if '-' in item:
            parts = item.split('-')
            if len(parts) != 2:
                raise ValidationError(f"{category}列范围格式错误：{item}")
            start, end = parts[0].strip(), parts[1].strip()
            start_idx = resolve_single_col(start, header_map)
            end_idx = resolve_single_col(end, header_map)
            if start_idx > max_col:
                raise ValidationError(f"{category}列 '{start}' 超出工作表列范围（共{max_col}列）")
            if end_idx > max_col:
                raise ValidationError(f"{category}列 '{end}' 超出工作表列范围（共{max_col}列）")
        else:
            idx = resolve_single_col(item, header_map)
            if idx > max_col:
                raise ValidationError(f"{category}列 '{item}' 超出工作表列范围（共{max_col}列）")

@subtotal_bp.route('/process', methods=['POST'])
@handle_exceptions()
def process():
    file = request.files.get('file')
    if not file or file.filename == '':
        raise ValidationError('请上传文件')

    original_filename = file.filename
    key_cols_raw = request.form.get('key_cols', '').strip()
    if not key_cols_raw:
        raise ValidationError('请指定主键列')

    dimension_cols_raw = request.form.get('dimension_cols', '').strip() or None
    calc_cols_raw = request.form.get('calc_cols', '').strip()
    if not calc_cols_raw:
        raise ValidationError('请指定计算列')

    agg_func_default = request.form.get('agg_func_default', 'sum').strip()
    agg_overrides_raw = request.form.get('agg_overrides', '').strip()
    sheets_mode = request.form.get('sheets_mode', '1').strip()

    file_service = current_app.container.file_service()
    path = file_service.save_upload(file, allowed_extensions={'xlsx', 'xls'})

    wb = openpyxl.load_workbook(path)
    first_ws = wb.active
    max_col = first_ws.max_column
    header_map = _build_header_map(first_ws)
    wb.close()

    _check_columns(key_cols_raw, '主键', header_map, max_col)
    _check_columns(calc_cols_raw, '计算', header_map, max_col)
    if dimension_cols_raw:
        _check_columns(dimension_cols_raw, '维度', header_map, max_col)

    agg_funcs = {'default': agg_func_default}
    if agg_overrides_raw:
        for part in agg_overrides_raw.split(','):
            if ':' in part:
                col, func = part.split(':', 1)
                col = col.strip().upper()
                func = func.strip().lower()
                if func not in ('sum', 'avg', 'min', 'max', 'count'):
                    raise ValidationError(f"不支持的聚合函数: {func}")
                try:
                    openpyxl.utils.column_index_from_string(col)
                except:
                    raise ValidationError(f"无效的列字母: {col}")
                agg_funcs[col] = func

    wb2 = openpyxl.load_workbook(path, read_only=True)
    total_rows = sum(sheet.max_row for sheet in wb2.worksheets)
    wb2.close()
    data_rows = total_rows - 1

    app = current_app._get_current_object()

    if data_rows > 5000:
        task_storage = app.container.task_storage()

        @async_task(original_filename='subtotal')
        def run_subtotal(task_id):
            with app.app_context():
                def progress_cb(percent, message):
                    task_storage.update(task_id, progress_percent=percent, progress_message=message)
                try:
                    progress_cb(0, "开始处理...")
                    engine = SubtotalEngine(
                        path, key_cols_raw, calc_cols_raw, agg_funcs,
                        dimension_cols=dimension_cols_raw,
                        sheets=sheets_mode,
                        original_filename=original_filename
                    )
                    out_path = engine.process()
                    progress_cb(100, "完成")
                    # 保存原始文件名到任务数据中，便于下载时生成统一名称
                    return {'result_path': out_path, 'original_filename': original_filename}
                except Exception as e:
                    import traceback
                    error_msg = f"{str(e)}\n{traceback.format_exc()}"
                    task_storage.append_log(task_id, error_msg)
                    raise
                finally:
                    file_service.cleanup([path])

        result = run_subtotal()
        return jsonify({'task_id': result['task_id']})
    else:
        try:
            engine = SubtotalEngine(
                path, key_cols_raw, calc_cols_raw, agg_funcs,
                dimension_cols=dimension_cols_raw,
                sheets=sheets_mode,
                original_filename=original_filename
            )
            out_path = engine.process()
            download_name = generate_output_filename('subtotal', original_filename, suffix='_小计')
            resp = send_file(out_path, as_attachment=True,
                             download_name=download_name,
                             mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            resp.headers['Content-Disposition'] = build_content_disposition(download_name)
            return resp
        finally:
            file_service.cleanup([path])

@subtotal_bp.route('/status/<task_id>')
def status(task_id):
    storage = current_app.container.task_storage()
    data = storage.load(task_id)
    if not data:
        return jsonify({'error': '任务不存在'}), 404
    log = storage.get_log(task_id, max_lines=100)
    data['log'] = log
    return jsonify(data)

@subtotal_bp.route('/download/<task_id>')
def download(task_id):
    storage = current_app.container.task_storage()
    data = storage.load(task_id)
    if not data or data.get('status') != 'completed':
        return jsonify({'error': '结果不可用'}), 400
    path = data.get('result_path')
    if not path or not os.path.exists(path):
        return jsonify({'error': '文件不存在'}), 404

    # 使用统一命名函数生成下载名
    original_filename = data.get('original_filename', 'data.xlsx')
    download_name = generate_output_filename('subtotal', original_filename, suffix='_小计')
    resp = send_file(path, as_attachment=True, download_name=download_name)
    resp.headers['Content-Disposition'] = build_content_disposition(download_name)
    return resp
