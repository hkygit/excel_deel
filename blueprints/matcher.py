# blueprints/matcher.py
import os
import io
import logging
import tempfile
import pandas as pd
from openpyxl import Workbook
from flask import Blueprint, render_template, request, send_file, current_app

from core.decorators import handle_exceptions
from core.exceptions import ValidationError
from core.file_naming import generate_output_filename
from core.csv_utils import split_fields
from core.excel_helpers import safe_delete_file
from services.excel.matcher import ExcelMatcher

logger = logging.getLogger(__name__)

match_bp = Blueprint('match', __name__, url_prefix='/match', template_folder='../templates/matcher')

ALLOWED_TARGET_EXT = {'.xlsx'}


def _check_target_extension(filename: str):
    ext = os.path.splitext(filename)[1].lower()
    if ext not in ALLOWED_TARGET_EXT:
        raise ValidationError(
            f"目标文件格式不支持，请使用 .xlsx 格式（当前: {ext}）。如为 .xls 文件，请先用 Excel 另存为 .xlsx。"
        )


def _fix_target_extension(filepath: str) -> str:
    if not filepath.lower().endswith('.xlsx'):
        new_path = filepath + '.xlsx'
        try:
            os.rename(filepath, new_path)
            return new_path
        except Exception:
            pass
    return filepath


def _parse_multi_cols(raw: str) -> list:
    if not raw:
        return []
    return split_fields(raw)


def _write_dataframe_to_worksheet(ws, df):
    """
    将 DataFrame 写入 worksheet（替代 pandas.to_excel，避免临时文件冲突）。
    """
    # 写表头
    ws.append([str(c) for c in df.columns])
    # 写数据
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))


def _ensure_excel_compatible(filepath: str, sheet_indices=None, file_label="文件") -> str:
    """
    只提取指定的 Sheet（默认为第一个），生成新的临时 .xlsx 文件。
    使用内存缓冲区保存，避免 Windows 临时文件问题。
    """
    logger.info(f"[规范化] 开始处理 {file_label}: {filepath}, 目标Sheet索引: {sheet_indices}")

    all_sheet_names = []
    try:
        import openpyxl as _oxl
        wb = _oxl.load_workbook(filepath, read_only=True, data_only=True)
        all_sheet_names = wb.sheetnames
        wb.close()
        logger.info(f"[规范化] 成功获取所有Sheet名称: {all_sheet_names}")
    except Exception as e:
        logger.warning(f"[规范化] 无法用 openpyxl 读取工作表名称: {e}，将使用 pandas 直接读取指定 Sheet")

    if sheet_indices is None:
        target_indices = [0]
    else:
        target_indices = [idx - 1 for idx in sheet_indices]

    dfs = {}
    failed_sheets = []

    for idx in target_indices:
        logger.info(f"[规范化] 尝试读取 {file_label} 的Sheet索引 {idx}（第{idx+1}个）")
        try:
            df = pd.read_excel(filepath, sheet_name=idx, dtype=str, engine='openpyxl')
            dfs[idx] = df
            logger.info(f"[规范化] {file_label} Sheet索引 {idx} 读取成功，行数: {len(df)}")
        except Exception as e:
            failed_sheets.append(f"第{idx+1}个Sheet")
            logger.warning(f"[规范化] 读取 {file_label} Sheet索引 {idx} 失败: {e}")

    if not dfs:
        raise ValidationError(
            f"{file_label} 无法读取所需的 Sheet（{', '.join(failed_sheets)}）。"
            "请确认这些 Sheet 未损坏且为有效的 Excel 表格。"
        )

    # === 使用 BytesIO 在内存中保存，避免 Windows 临时文件问题 ===
    wb_new = Workbook()
    default_sheet = wb_new.active
    wb_new.remove(default_sheet)

    for idx, df in dfs.items():
        try:
            sheet_name = all_sheet_names[idx] if idx < len(all_sheet_names) else f"Sheet{idx+1}"
        except Exception:
            sheet_name = f"Sheet{idx+1}"
        sheet_name = str(sheet_name).strip()[:31] or f"Sheet{idx+1}"
        for ch in r'[]:*?/\\':
            sheet_name = sheet_name.replace(ch, '_')

        ws = wb_new.create_sheet(title=sheet_name)
        _write_dataframe_to_worksheet(ws, df)
        logger.info(f"[规范化] 已写入 {file_label} Sheet: {sheet_name}")

    # 保存到内存缓冲区
    buffer = io.BytesIO()
    wb_new.save(buffer)
    wb_new.close()
    buffer.seek(0)

    # 将内存中的 Excel 写入一个安全的磁盘文件（供后续 openpyxl 处理）
    tmp_dir = tempfile.mkdtemp(prefix='matcher_norm_')
    new_path = os.path.join(tmp_dir, 'normalized.xlsx')
    with open(new_path, 'wb') as f:
        f.write(buffer.getvalue())

    logger.info(f"[规范化] {file_label} 新临时文件生成: {new_path}")
    return new_path


@match_bp.route('/')
def index():
    return render_template('matcher/index.html')


@match_bp.route('/process', methods=['POST'])
@handle_exceptions()
def process():
    logger.info("[匹配填充] 收到 POST 请求")
    file_service = current_app.container.file_service()

    target_file = request.files.get('target_file')
    source_file = request.files.get('source_file')
    if not target_file or not source_file:
        raise ValidationError('请上传两个文件')

    _check_target_extension(target_file.filename)

    logger.info("[匹配填充] 保存上传文件...")
    target_raw_path = file_service.save_upload(target_file)
    source_raw_path = file_service.save_upload(source_file)

    target_raw_path = _fix_target_extension(target_raw_path)

    source_sheets_raw = request.form.get('source_sheets', '').strip()
    sheet_indices = None
    if source_sheets_raw:
        try:
            sheet_indices = [int(x.strip()) for x in source_sheets_raw.split(',') if x.strip()]
            if not sheet_indices:
                sheet_indices = None
        except ValueError:
            raise ValidationError('Sheet索引格式错误，请输入逗号分隔的数字，例如 1,2,3')

    temp_files_to_clean = []

    try:
        logger.info("[匹配填充] 开始规范化目标文件（只读第1个Sheet）")
        target_path = _ensure_excel_compatible(target_raw_path, sheet_indices=[1], file_label="目标文件")
        temp_files_to_clean.append(target_path)
        temp_files_to_clean.append(target_raw_path)

        logger.info("[匹配填充] 开始规范化源文件")
        source_path = _ensure_excel_compatible(source_raw_path, sheet_indices=sheet_indices, file_label="源文件")
        temp_files_to_clean.append(source_path)
        temp_files_to_clean.append(source_raw_path)
    except ValidationError as e:
        for p in temp_files_to_clean:
            safe_delete_file(p)
        raise

    target_key_cols_raw = request.form.get('target_key_cols', '').strip()
    source_match_cols_raw = request.form.get('source_match_cols', '').strip()
    source_value_cols_raw = request.form.get('source_value_col', '').strip()
    result_cols_raw = request.form.get('result_col', '').strip()
    allow_duplicate_keys = request.form.get('allow_duplicate_keys') == 'on'

    if not target_key_cols_raw or not source_match_cols_raw or not source_value_cols_raw or not result_cols_raw:
        raise ValidationError('请填写完整配置项')

    target_key_cols = _parse_multi_cols(target_key_cols_raw)
    source_match_cols = _parse_multi_cols(source_match_cols_raw)
    source_value_cols = _parse_multi_cols(source_value_cols_raw)
    result_cols = _parse_multi_cols(result_cols_raw)

    if len(source_value_cols) != len(result_cols):
        raise ValidationError(f'源文件值列数量({len(source_value_cols)})与结果列数量({len(result_cols)})不一致')
    if not target_key_cols or not source_match_cols:
        raise ValidationError('主键列和匹配列不能为空')
    if len(target_key_cols) != len(source_match_cols):
        raise ValidationError(
            f'目标主键列数量({len(target_key_cols)})与源匹配列数量({len(source_match_cols)})不一致'
        )

    config = {
        'target_key_cols': target_key_cols,
        'source_match_cols': source_match_cols,
        'source_value_cols': source_value_cols,
        'result_cols': result_cols,
        'allow_duplicate_keys': allow_duplicate_keys
    }

    logger.info("[匹配填充] 开始执行 ExcelMatcher.match()")
    matcher = ExcelMatcher(target_path, source_path, config, sheet_indices)
    success, result = matcher.match()
    if not success:
        error_msg = result if isinstance(result, str) else '\n'.join(result)
        raise ValidationError(error_msg)

    for p in temp_files_to_clean:
        if p and os.path.exists(p) and p != result:
            safe_delete_file(p)

    download_name = generate_output_filename('match', target_file.filename, suffix='_匹配结果')
    logger.info("[匹配填充] 处理完成，返回文件")
    return send_file(result, as_attachment=True, download_name=download_name)


@match_bp.route('/debug', methods=['POST'])
@handle_exceptions()
def debug():
    file_service = current_app.container.file_service()

    target_file = request.files.get('target_file')
    source_file = request.files.get('source_file')
    if not target_file or not source_file:
        return render_template('error.html', error='请上传两个文件'), 400

    _check_target_extension(target_file.filename)

    target_raw_path = file_service.save_upload(target_file)
    target_raw_path = _fix_target_extension(target_raw_path)
    source_raw_path = file_service.save_upload(source_file)

    source_sheets_raw = request.form.get('source_sheets', '').strip()
    sheet_indices = None
    if source_sheets_raw:
        try:
            sheet_indices = [int(x.strip()) for x in source_sheets_raw.split(',') if x.strip()]
            if not sheet_indices:
                sheet_indices = None
        except ValueError:
            return render_template('error.html', error='Sheet索引格式错误，请输入逗号分隔的数字'), 400

    temp_files_to_clean = []

    try:
        target_path = _ensure_excel_compatible(target_raw_path, sheet_indices=[1], file_label="目标文件")
        temp_files_to_clean.append(target_path)
        temp_files_to_clean.append(target_raw_path)

        source_path = _ensure_excel_compatible(source_raw_path, sheet_indices=sheet_indices, file_label="源文件")
        temp_files_to_clean.append(source_path)
        temp_files_to_clean.append(source_raw_path)
    except ValidationError as e:
        for p in temp_files_to_clean:
            safe_delete_file(p)
        raise

    target_key_cols_raw = request.form.get('target_key_cols', '').strip()
    source_match_cols_raw = request.form.get('source_match_cols', '').strip()
    source_value_cols_raw = request.form.get('source_value_col', '').strip()
    result_cols_raw = request.form.get('result_col', '').strip()
    allow_duplicate_keys = request.form.get('allow_duplicate_keys') == 'on'

    if not target_key_cols_raw or not source_match_cols_raw or not source_value_cols_raw or not result_cols_raw:
        return render_template('error.html', error='请填写完整配置项'), 400

    target_key_cols = _parse_multi_cols(target_key_cols_raw)
    source_match_cols = _parse_multi_cols(source_match_cols_raw)
    source_value_cols = _parse_multi_cols(source_value_cols_raw)
    result_cols = _parse_multi_cols(result_cols_raw)

    if len(source_value_cols) != len(result_cols):
        return render_template(
            'error.html',
            error=f'源文件值列数量({len(source_value_cols)})与结果列数量({len(result_cols)})不一致'
        ), 400

    config = {
        'target_key_cols': target_key_cols,
        'source_match_cols': source_match_cols,
        'source_value_cols': source_value_cols,
        'result_cols': result_cols,
        'allow_duplicate_keys': allow_duplicate_keys
    }

    matcher = ExcelMatcher(target_path, source_path, config, sheet_indices)
    info = matcher.debug_info()
    if 'error' in info:
        return render_template('error.html', error=info['error']), 400

    for p in temp_files_to_clean:
        safe_delete_file(p)

    return render_template('debug.html', config=config, info=info)